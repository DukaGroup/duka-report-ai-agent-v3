"""Read back every protected source cell and verify generated calculation caches."""

import io
import math
import re
import zipfile
from datetime import datetime
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.cell.text import Text

from .currencies import display_currency
from .workbooks import ReportDataError, excel_number, current_workbook
from .workbook_contract import datasets, metric, comparison, select_currency
from .workbook_locale import presentation


def equivalent(actual, expected):
    if expected is None or expected == "":
        return actual in (None, "")
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual == expected
    if isinstance(expected, (int, float)):
        return isinstance(actual, (int, float)) and not isinstance(actual, bool) and math.isclose(actual, expected, rel_tol=2e-15, abs_tol=1e-10)
    if isinstance(expected, datetime):
        return isinstance(actual, datetime) and abs((actual-expected).total_seconds()) < 0.001
    return actual == expected


def source_value(data, row, field, for_output=False):
    value = row.get(field)
    if field in data["currency_columns"]:
        value = display_currency(value)
    kind = data["types"].get(field)
    if value is None or isinstance(value, bool):
        return value
    if kind in ("decimal", "numeric", "money", "smallmoney", "number", "integer", "float", "real", "int", "bigint", "smallint", "tinyint") or isinstance(value, (int, float)):
        converted = excel_number(value, True)
        return converted if converted is not None else str(value)
    if not for_output and kind in ("date", "datetime", "datetime2", "smalldatetime", "datetimeoffset"):
        try:
            return datetime.fromisoformat(str(value)).replace(tzinfo=None)
        except ValueError:
            pass
    return str(value)


def restore_ooxml_strings(workbook, content):
    """Read OOXML escapes once, preserving literal strings such as _x000D_.

    openpyxl 3.1.5 leaves carriage-return escapes undecoded and strips the
    literal-underscore escape. Restore its shared string list from the XML
    before read-only cells are streamed; do not weaken source comparisons.
    """
    if not workbook.worksheets:
        return
    strings = workbook.worksheets[0]._shared_strings
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if 'xl/sharedStrings.xml' not in archive.namelist():
            return
        index = 0
        with archive.open('xl/sharedStrings.xml') as source:
            for _, node in ElementTree.iterparse(source, events=('end',)):
                if node.tag != '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}si':
                    continue
                value = Text.from_tree(node).content
                value = re.sub(r'_x([0-9a-fA-F]{4})_', lambda match: chr(int(match[1],16)), value)
                # OOXML may encode a supplementary Unicode character as a pair.
                value = value.encode('utf-16',errors='surrogatepass').decode('utf-16')
                strings[index] = value
                index += 1
                node.clear()


def validate_workbook(content, document, snapshot, checks, custom=False):
    data = datasets(document, snapshot, source_only=custom)
    _, labels, _ = presentation(document)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if len(archive.infolist()) > 10000 or sum(item.file_size for item in archive.infolist()) > 2 * 1024 * 1024 * 1024:
            raise ReportDataError("Generated workbook exceeds validation capacity.")
        if any("externalLinks" in name or "vbaProject" in name or "connections.xml" in name for name in archive.namelist()):
            raise ReportDataError("Generated workbook contains an unsupported external connection.")
        tables = {}
        for name in archive.namelist():
            if name.startswith("xl/tables/") and name.endswith(".xml"):
                table = ElementTree.fromstring(archive.read(name))
                tables[table.get("name")] = table.get("ref")
    if not current_workbook(io.BytesIO(content)):
        raise ReportDataError("Generated workbook version is invalid.")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        restore_ooxml_strings(workbook, content)
        from .balance_workbook import validate_balance
        validate_balance(workbook, snapshot, labels, equivalent)
        for source in data:
            if source["sheet_name"] not in workbook.sheetnames:
                raise ReportDataError("Generated workbook is missing a protected source sheet.")
            sheet = workbook[source["sheet_name"]]
            if sheet.sheet_state != "visible":
                raise ReportDataError("Protected source sheets must remain visible.")
            if not source["rows"]:
                continue
            if source["table_name"] not in tables:
                raise ReportDataError("Generated workbook is missing a complete source table.")
            first_column, header, _, last_row = range_boundaries(tables[source["table_name"]])
            if first_column != 1 or header != source["first_row"]-1 or last_row != header+len(source["rows"]):
                raise ReportDataError("Generated workbook source row count or alignment does not match the snapshot.")
            headings = next(sheet.iter_rows(min_row=header, max_row=header, max_col=len(source["fields"]), values_only=True))
            if list(headings) != [source["column_labels"].get(field, field) for field in source["fields"]]:
                raise ReportDataError("Generated workbook source column headings are misaligned.")
            for source_row, actual in zip(source["rows"], sheet.iter_rows(min_row=header+1, max_row=last_row, max_col=len(source["fields"]), values_only=True), strict=True):
                if any(not equivalent(value, source_value(source, source_row, field)) for field, value in zip(source["fields"], actual, strict=True)):
                    raise ReportDataError("Generated workbook source values or currencies do not match the snapshot.")
        for check in checks:
            sheet = workbook[check["sheet"]]
            item = check["request"]
            if check["kind"] == "table":
                source = data[item["dataset"]]
                rows = select_currency(source, item["currency"], item.get("currency_column"))[0] if item.get("currency") else source["rows"]
                fields = item.get("columns") or source["fields"]
                if check["row_count"] != len(rows):
                    raise ReportDataError("Generated output table is incomplete.")
                if rows:
                    for source_row, actual in zip(rows, sheet.iter_rows(min_row=check["row"], max_row=check["row"]+len(rows)-1, min_col=check.get("column",1), max_col=check.get("column",1)+len(fields)-1, values_only=True), strict=True):
                        if any(not equivalent(value, source_value(source, source_row, field, True)) for field, value in zip(fields, actual, strict=True)):
                            raise ReportDataError("Generated output table rows are misaligned.")
            else:
                value = metric(data[item["dataset"]], item["column"], item.get("currency"), item["operation"], item.get("currency_column"))[0] if check["kind"] == "metric" else comparison(data, item)[0]
                expected = labels["not_available"] if value is None else excel_number(value, True) if excel_number(value, True) is not None else str(value)
                actual = next(sheet.iter_rows(min_row=check["row"], max_row=check["row"], min_col=check["column"], max_col=check["column"], values_only=True))[0]
                if not equivalent(actual, expected):
                    raise ReportDataError("Generated workbook calculation does not reconcile with the source snapshot.")
    finally:
        workbook.close()
    # Validate formula safety separately from cached results.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                with archive.open(name) as stream:
                    for _, element in ElementTree.iterparse(stream, events=("end",)):
                        if element.tag.endswith("}f") and re.search(r"\[|WEBSERVICE|HYPERLINK|DDE|RTD|FILTERXML|INDIRECT", element.text or "", re.I):
                            raise ReportDataError("Generated workbook contains an unsupported formula.")
                        element.clear()
    return {"source_rows": [len(source["rows"]) for source in data], "verified_calculations": sum(check["kind"] != "table" for check in checks)}
