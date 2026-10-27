"""Bounded, read-only previews of saved report workbooks."""

from datetime import date, datetime, time
from math import isfinite
from re import fullmatch

from openpyxl import load_workbook

from .workbook_visuals import sheet_visuals


PREVIEW_ROWS = 50
PREVIEW_COLUMNS = 16


class PreviewSheetNotFound(ValueError):
    pass


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and not isfinite(value):
        return ""
    return str(value)


def display_text(cell):
    """Apply the numeric formats used by generated reports without changing values."""
    value = cell.value
    raw = cell_text(value)
    if not isinstance(value, (int, float)) or isinstance(value, bool) or not isfinite(value):
        return raw
    sections = cell.number_format.split(";")
    section = sections[1] if value < 0 and len(sections) > 1 else sections[0]
    if value == 0 and len(sections) > 2 and '"—"' in sections[2]:
        return "—"
    if not fullmatch(r'[0#?,.%()\[\]A-Za-z\s"-]+', section) or "E+" in section.upper():
        return raw
    if "0" not in section and "#" not in section:
        return raw
    percentage = "%" in section
    decimal_pattern = section.split(".", 1)[1].split("%", 1)[0] if "." in section else ""
    fixed = decimal_pattern.count("0")
    optional = decimal_pattern.count("#")
    if fixed + optional > 6:
        return raw
    amount = abs(value) * (100 if percentage else 1)
    rendered = f"{amount:,.{fixed + optional}f}" if "," in section else f"{amount:.{fixed + optional}f}"
    if optional and "." in rendered:
        whole, fraction = rendered.split(".")
        fraction = fraction.rstrip("0")
        rendered = whole + ("." + fraction if fraction else ("." + "0" * fixed if fixed else ""))
    if value < 0:
        rendered = f"({rendered})" if "(" in section else f"-{rendered}"
    return rendered + ("%" if percentage else "")


def rgb_color(color):
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    value = color.rgb[-6:]
    return "#" + value.upper() if fullmatch(r"[0-9A-Fa-f]{6}", value) else None


def preview_cell(cell):
    value = cell.value
    font = cell.font
    fill = cell.fill
    alignment = cell.alignment.horizontal if cell.alignment else None
    return {
        "text": display_text(cell),
        "raw": cell_text(value),
        "bold": bool(font.bold) if value is not None and font else False,
        "italic": bool(font.italic) if value is not None and font else False,
        "numeric": isinstance(value, (int, float)) and not isinstance(value, bool),
        "color": rgb_color(font.color) if value is not None and font else None,
        "background": rgb_color(fill.fgColor) if fill and fill.patternType == "solid" else None,
        "align": alignment if alignment in ("left", "center", "right") else None,
        "wrap": bool(cell.alignment.wrap_text) if cell.alignment else False,
    }


def workbook_preview(path, sheet_index=0, row_start=1, column_start=1):
    """Read one bounded window. Every row and column remains reachable by paging."""
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheets = [
            {
                "name": sheet.title,
                "row_count": sheet.max_row or 0,
                "column_count": sheet.max_column or 0,
            }
            for sheet in workbook.worksheets
        ]
        if sheet_index >= len(sheets):
            raise PreviewSheetNotFound("Fleta e kërkuar nuk u gjet.")
        sheet = workbook.worksheets[sheet_index]
        row_count = sheets[sheet_index]["row_count"]
        column_count = sheets[sheet_index]["column_count"]
        row_start = min(max(1, row_start), max(1, row_count))
        column_start = min(max(1, column_start), max(1, column_count))
        visible_rows = min(PREVIEW_ROWS, max(0, row_count - row_start + 1))
        visible_columns = min(PREVIEW_COLUMNS, max(0, column_count - column_start + 1))
        rows = []
        if visible_rows and visible_columns:
            for row in sheet.iter_rows(
                min_row=row_start,
                max_row=row_start + visible_rows - 1,
                min_col=column_start,
                max_col=column_start + visible_columns - 1,
            ):
                rows.append([preview_cell(cell) for cell in row])
        merged_ranges, charts = sheet_visuals(str(path), path.stat().st_mtime_ns, sheet_index)
        end_row = row_start + visible_rows - 1
        end_column = column_start + visible_columns - 1
        return {
            "sheets": sheets,
            "sheet_index": sheet_index,
            "row_start": row_start,
            "column_start": column_start,
            "rows": rows,
            "visible_rows": visible_rows,
            "visible_columns": visible_columns,
            "has_more_rows": end_row < row_count,
            "has_more_columns": end_column < column_count,
            "merged": [
                {"start_row": min_row, "start_column": min_col, "end_row": max_row, "end_column": max_col}
                for min_col, min_row, max_col, max_row in merged_ranges
                if min_row <= end_row and max_row >= row_start and min_col <= end_column and max_col >= column_start
            ],
            "charts": charts,
        }
    finally:
        workbook.close()
