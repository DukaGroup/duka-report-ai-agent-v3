"""Isolated, resource-limited extraction. Uploaded formulas/macros are never run."""

import csv
import io
import json
import math
import re
import sys
import zipfile
import posixpath
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_ROWS, MAX_COLS, MAX_CELLS, MAX_TEXT = 50000, 100, 2500000, 200000
MAX_ARCHIVE_BYTES = 128 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_RATIO = 250
MAX_PARSED_BYTES = 96 * 1024 * 1024
STREAM_ARCHIVE_BYTES = 24 * 1024 * 1024
STREAM_CELL_COUNT = 200000


def scalar(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)) or value is None:
        if isinstance(value, str) and len(value) > 5000:
            raise ValueError("An uploaded cell exceeds the text limit.")
        return value
    return str(value)


def table(identifier, name, matrix):
    if not matrix:
        return None
    if len(matrix) > MAX_ROWS + 1 or max(map(len, matrix)) > MAX_COLS:
        raise ValueError("Uploaded tables support at most 50000 rows and 100 columns.")
    headers, used = [], set()
    for i, cell in enumerate(matrix[0]):
        base = re.sub(r"[\x00-\x1f]", " ", str(cell or f"Column {i + 1}"))[:90].strip() or f"Column {i + 1}"
        label, suffix = base, 1
        while label.casefold() in used:
            suffix += 1
            label = f"{base}_{suffix}"
        used.add(label.casefold())
        headers.append(label)
    rows = [
        {key: scalar(row[i]) if i < len(row) else None for i, key in enumerate(headers)}
        for row in matrix[1:]
        if any(v is not None and v != "" for v in row)
    ]
    columns = []
    for key in headers:
        values = [r[key] for r in rows if r[key] is not None and r[key] != ""]
        kind = (
            "number"
            if values and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in values)
            else "text"
        )
        if (
            kind == "text"
            and values
            and all(isinstance(v, str) and re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", v.strip()) for v in values)
        ):
            kind = "decimal"
        if values and all(isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T.*)?", v) for v in values):
            kind = "date"
        columns.append({"name": key, "type": kind})
    return {"id": identifier, "name": name, "columns": columns, "rows": rows, "row_count": len(rows)}


def guard_zip(path, *, max_uncompressed=MAX_ARCHIVE_BYTES):
    with zipfile.ZipFile(path) as z:
        members = z.infolist()
        total_size = sum(i.file_size for i in members)
        if (
            len(members) > 3000
            or total_size > max_uncompressed
            or any(i.file_size > MAX_ARCHIVE_MEMBER_BYTES for i in members)
            or any(
                i.file_size > 1024 * 1024
                and i.file_size / max(i.compress_size, 1) > MAX_ARCHIVE_RATIO
                for i in members
            )
        ):
            raise ValueError("The uploaded document is too large after decompression.")
        if any(i.flag_bits & 1 for i in members):
            raise ValueError("Encrypted uploads are not supported.")
        return total_size


def _xlsx_sheet_parts(archive):
    """Map workbook sheet order to worksheet XML paths and associated table refs."""
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets = {
        rel.attrib["Id"]: rel.attrib["Target"]
        for rel in relationships
        if rel.attrib.get("Target")
    }
    rel_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    sheets = []
    for sheet in workbook.findall(".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
        target = rel_targets.get(sheet.attrib.get(rel_id))
        if not target:
            raise ValueError("The uploaded Excel workbook has an invalid sheet reference.")
        sheet_path = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl", target))
        )
        table_specs = []
        rel_path = posixpath.join(
            posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels"
        )
        if rel_path in archive.namelist():
            sheet_rels = ET.fromstring(archive.read(rel_path))
            for rel in sheet_rels:
                if not rel.attrib.get("Type", "").endswith("/table"):
                    continue
                rel_target = rel.attrib.get("Target", "")
                table_path = (
                    rel_target.lstrip("/")
                    if rel_target.startswith("/")
                    else posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), rel_target))
                )
                table_xml = ET.fromstring(archive.read(table_path))
                columns = [
                    item.attrib.get("name", "")
                    for item in table_xml.findall(
                        ".//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}tableColumn"
                    )
                ]
                style = table_xml.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}tableStyleInfo")
                table_specs.append(
                    {
                        "ref": table_xml.attrib.get("ref", ""),
                        "name": table_xml.attrib.get("displayName") or table_xml.attrib.get("name", "Table"),
                        "columns": columns,
                        "style": style.attrib.get("name") if style is not None else None,
                        "totals": int(table_xml.attrib.get("totalsRowCount", "0") or 0),
                    }
                )
        sheets.append({"name": sheet.attrib.get("name", "Sheet"), "path": sheet_path, "tables": table_specs})
    return sheets


def _clean_headers(values):
    headers, used = [], set()
    for i, cell in enumerate(values):
        base = re.sub(r"[\x00-\x1f]", " ", str(cell or f"Column {i + 1}"))[:90].strip() or f"Column {i + 1}"
        label, suffix = base, 1
        while label.casefold() in used:
            suffix += 1
            label = f"{base}_{suffix}"
        used.add(label.casefold())
        headers.append(label)
    return headers


def _stream_table(identifier, name, headers, rows):
    """Retain large tables as compact row arrays; consumers map them to column names."""
    headers = _clean_headers(headers)
    if len(headers) > MAX_COLS:
        raise ValueError("Uploaded tables support at most 50000 rows and 100 columns.")
    values = []
    has_values = [False] * len(headers)
    all_numeric = [True] * len(headers)
    all_decimal_text = [True] * len(headers)
    all_dates = [True] * len(headers)
    for row in rows:
        row = list(row)
        if len(row) < len(headers):
            row.extend([None] * (len(headers) - len(row)))
        row = [scalar(v) for v in row[: len(headers)]]
        if not any(v is not None and v != "" for v in row):
            continue
        values.append(row)
        if len(values) > MAX_ROWS:
            raise ValueError("Uploaded tables support at most 50000 rows and 100 columns.")
        for index, value in enumerate(row):
            if value is None or value == "":
                continue
            has_values[index] = True
            all_numeric[index] &= isinstance(value, (int, float)) and not isinstance(value, bool)
            all_decimal_text[index] &= isinstance(value, str) and bool(
                re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", value.strip())
            )
            all_dates[index] &= isinstance(value, str) and bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T.*)?", value))
    columns = []
    for index, header in enumerate(headers):
        kind = "number" if has_values[index] and all_numeric[index] else "text"
        if kind == "text" and has_values[index] and all_decimal_text[index]:
            kind = "decimal"
        if has_values[index] and all_dates[index]:
            kind = "date"
        columns.append({"name": header, "type": kind})
    return {
        "id": identifier,
        "name": name,
        "columns": columns,
        "rows": values,
        "row_encoding": "array",
        "row_count": len(values),
    }


def _preview_sheet(ws, parts, remaining):
    """Capture a bounded text preview; full Excel table rows are stored separately."""
    if remaining <= 0:
        return 0
    label = f"[{ws.title}]"
    parts.append(label)
    remaining -= len(label) + 1
    for row in ws.iter_rows(values_only=True):
        content = " | ".join(str(scalar(cell)) for cell in row if cell is not None)
        if not content:
            continue
        if len(content) > remaining:
            parts.append(content[:remaining])
            return 0
        parts.append(content)
        remaining -= len(content) + 1
        if remaining <= 0:
            return 0
    return remaining


def extract_xlsx_streaming(path):
    """Read a bounded, larger workbook without materializing worksheet cells in memory."""
    from openpyxl import load_workbook
    from openpyxl.utils.cell import range_boundaries

    with zipfile.ZipFile(path) as archive:
        sheets = _xlsx_sheet_parts(archive)
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if len(wb.worksheets) > 20:
            raise ValueError("Upload at most 20 worksheets per file.")
        if len(sheets) != len(wb.worksheets):
            raise ValueError("The uploaded Excel workbook has inconsistent sheet metadata.")
        tables, parts, warnings, layout = [], [], [], {}
        text_remaining = MAX_TEXT
        total_cells = 0
        for si, (ws, spec) in enumerate(zip(wb.worksheets, sheets)):
            max_row, max_col = ws.max_row or 0, ws.max_column or 0
            if max_row > MAX_ROWS + 50 or max_col > MAX_COLS:
                raise ValueError("The uploaded worksheet exceeds the row or column limit.")
            total_cells += max_row * max_col
            if total_cells > MAX_CELLS:
                raise ValueError("The upload exceeds the cell limit.")
            found = False
            for ti, table_spec in enumerate(spec["tables"]):
                try:
                    min_col, min_row, max_col, max_row = range_boundaries(table_spec["ref"])
                except Exception as exc:
                    raise ValueError("The uploaded Excel table has an invalid range.") from exc
                if min_row < 1 or min_col < 1 or max_row - min_row > MAX_ROWS or max_col - min_col + 1 > MAX_COLS:
                    raise ValueError("Uploaded tables support at most 50000 rows and 100 columns.")
                data_end = max_row - min(table_spec["totals"], max_row - min_row)
                row_iter = ws.iter_rows(
                    min_row=min_row, max_row=data_end, min_col=min_col, max_col=max_col, values_only=True
                )
                header = next(row_iter, None)
                if header is None:
                    continue
                item = _stream_table(
                    f"s{si + 1}_t{ti + 1}",
                    ws.title,
                    header,
                    row_iter,
                )
                tables.append(item)
                found = True
                if not layout:
                    layout = {
                        "sheet_names": [sheet.title for sheet in wb.worksheets],
                        "font_name": "Calibri",
                        "header_color": "#205C49",
                        "header_font_color": "#FFFFFF",
                        "table_style": table_spec["style"] or "Table Style Medium 4",
                        "column_widths": [18] * len(item["columns"]),
                        "header_row": min(15, max(4, min_row - 1)),
                        "chart_type": "column",
                        "columns": [column["name"] for column in item["columns"]],
                        "landscape": True,
                        "layout_version": 1,
                        "streaming_limited": True,
                    }
            if not found:
                first = None
                for index, row in enumerate(
                    ws.iter_rows(min_row=1, max_row=min(30, ws.max_row or 0), max_col=min(MAX_COLS, ws.max_column or 0), values_only=True)
                ):
                    if sum(isinstance(value, str) and bool(value.strip()) for value in row) >= 2:
                        first = index + 1
                        break
                if first is not None:
                    headers = next(
                        ws.iter_rows(min_row=first, max_row=first, max_col=min(MAX_COLS, ws.max_column or 0), values_only=True)
                    )
                    rows = ws.iter_rows(
                        min_row=first + 1,
                        max_row=ws.max_row,
                        max_col=min(MAX_COLS, ws.max_column or 0),
                        values_only=True,
                    )
                    item = _stream_table(f"s{si + 1}", ws.title, headers, rows)
                    if item["rows"]:
                        tables.append(item)
                        if not layout:
                            layout = {
                                "sheet_names": [sheet.title for sheet in wb.worksheets],
                                "font_name": "Calibri",
                                "header_color": "#205C49",
                                "header_font_color": "#FFFFFF",
                                "table_style": "Table Style Medium 4",
                                "column_widths": [18] * len(item["columns"]),
                                "header_row": min(15, max(4, first - 1)),
                                "chart_type": "column",
                                "columns": [column["name"] for column in item["columns"]],
                                "landscape": True,
                                "layout_version": 1,
                                "streaming_limited": True,
                            }
            if text_remaining > 0:
                text_remaining = _preview_sheet(ws, parts, text_remaining)
        text = "\n".join(parts)
        warnings.append(
            "Only cached Excel formula results are read; formulas and macros are never executed. Missing caches remain missing data."
        )
        warnings.append(
            "This large workbook was read in bounded streaming mode. Table data is complete; the text preview and captured formatting are limited."
        )
        return {
            "tables": tables,
            "text": text[:MAX_TEXT],
            "text_truncated": text_remaining <= 0,
            "warnings": warnings,
            "format": layout,
        }
    finally:
        wb.close()


def _xlsx_requires_streaming(path, archive_size):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if len(wb.worksheets) > 20:
            raise ValueError("Upload at most 20 worksheets per file.")
        cells = 0
        for ws in wb.worksheets:
            if (ws.max_row or 0) > MAX_ROWS + 50 or (ws.max_column or 0) > MAX_COLS:
                raise ValueError("The uploaded worksheet exceeds the row or column limit.")
            cells += (ws.max_row or 0) * (ws.max_column or 0)
            if cells > MAX_CELLS:
                raise ValueError("The upload exceeds the cell limit.")
        return archive_size >= STREAM_ARCHIVE_BYTES or cells > STREAM_CELL_COUNT
    finally:
        wb.close()


def extract(path, extension):
    tables, parts, warnings, layout = [], [], [], {}
    archive_size = 0
    if extension in (".xlsx", ".docx"):
        archive_size = guard_zip(
            path,
            max_uncompressed=MAX_ARCHIVE_BYTES if extension == ".xlsx" else 50 * 1024 * 1024,
        )
    if extension == ".xlsx":
        if _xlsx_requires_streaming(path, archive_size):
            return extract_xlsx_streaming(path)
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries, get_column_letter

        wb = load_workbook(path, data_only=True, keep_links=False)
        if len(wb.worksheets) > 20:
            raise ValueError("Upload at most 20 worksheets per file.")
        cells = 0
        for si, ws in enumerate(wb.worksheets):
            if ws.max_row > MAX_ROWS + 50 or ws.max_column > MAX_COLS:
                raise ValueError("The uploaded worksheet exceeds the row or column limit.")
            cells += ws.max_row * ws.max_column
            if cells > MAX_CELLS:
                raise ValueError("The upload exceeds the cell limit.")
            parts.append(f"[{ws.title}]")
            for row in ws.iter_rows():
                content = " | ".join(str(scalar(c.value)) for c in row if c.value is not None)
                if content:
                    parts.append(content)
            found = False
            for ti, ref in enumerate(ws.tables.values()):
                a, b, c, d = range_boundaries(ref.ref)
                if ref.totalsRowCount:
                    d -= ref.totalsRowCount
                matrix = [
                    [cell.value for cell in row] for row in ws.iter_rows(min_row=b, max_row=d, min_col=a, max_col=c)
                ]
                item = table(f"s{si + 1}_t{ti + 1}", ws.title, matrix)
                if item:
                    tables.append(item)
                    found = True
                if not layout:
                    header = ws.cell(b, a)
                    color = header.fill.fgColor.rgb if header.fill.fgColor.type == "rgb" else None
                    font_color = (
                        header.font.color.rgb if header.font.color and header.font.color.type == "rgb" else None
                    )
                    chart = ws._charts[0] if ws._charts else next((c for s in wb.worksheets for c in s._charts), None)
                    chart_type = (
                        "line"
                        if chart and chart.__class__.__name__ == "LineChart"
                        else "bar"
                        if chart and getattr(chart, "type", None) == "bar"
                        else "column"
                    )
                    layout = {
                        "sheet_names": wb.sheetnames,
                        "font_name": header.font.name or "Calibri",
                        "header_color": "#" + color[-6:] if isinstance(color, str) else "#205C49",
                        "header_font_color": "#" + font_color[-6:] if isinstance(font_color, str) else "#FFFFFF",
                        "table_style": ref.tableStyleInfo.name if ref.tableStyleInfo else "Table Style Medium 4",
                        "column_widths": [
                            min(60, max(12, ws.column_dimensions[get_column_letter(i)].width or 18))
                            for i in range(a, c + 1)
                        ],
                        "header_row": min(15, max(4, b - 1)),
                        "chart_type": chart_type,
                        "columns": [str(v or "") for v in matrix[0]],
                        "landscape": ws.page_setup.orientation == "landscape",
                    }
            if not found:
                matrix = [[c.value for c in row] for row in ws.iter_rows()]
                first = next(
                    (
                        i
                        for i, row in enumerate(matrix[:30])
                        if sum(isinstance(v, str) and bool(v.strip()) for v in row) >= 2
                    ),
                    None,
                )
                if first is not None:
                    item = table(f"s{si + 1}", ws.title, matrix[first:])
                    if item:
                        tables.append(item)
                        if not layout:
                            header = ws.cell(first + 1, 1)
                            color = header.fill.fgColor.rgb if header.fill.fgColor.type == "rgb" else None
                            layout = {
                                "sheet_names": wb.sheetnames,
                                "font_name": header.font.name or "Calibri",
                                "header_color": "#" + color[-6:]
                                if isinstance(color, str) and color[-6:] != "000000"
                                else "#205C49",
                                "header_font_color": "#FFFFFF",
                                "header_row": min(15, max(4, first)),
                                "columns": [str(v or "") for v in matrix[first]],
                                "column_widths": [
                                    min(
                                        60,
                                        max(
                                            12,
                                            ws.column_dimensions[get_column_letter(i + 1)].width or 18,
                                        ),
                                    )
                                    for i in range(len(matrix[first]))
                                ],
                                "chart_type": "column",
                                "table_style": "Table Style Medium 4",
                            }
        from .reference_layout import workbook_layout
        layout.update(workbook_layout(wb))
        wb.close()
        warnings.append(
            "Only cached Excel formula results are read; formulas and macros are never executed. Missing caches remain missing data."
        )
    elif extension == ".csv":
        content = Path(path).read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
        matrix = list(csv.reader(io.StringIO(content), dialect))
        if len(matrix) > MAX_ROWS + 1:
            raise ValueError("Uploaded tables support at most 50000 rows and 100 columns.")
        for row in matrix[1:]:
            for i, value in enumerate(row):
                if (
                    re.fullmatch(r"-?(?:0|[1-9]\d*)(?:\.\d+)?", value.strip())
                    and len(value.replace(".", "").replace("-", "")) <= 15
                ):
                    row[i] = float(value) if "." in value else int(value)
        item = table("csv", "CSV", matrix)
        if item:
            tables.append(item)
        warnings.append(
            "CSV numeric detection uses a decimal point; ambiguous comma-decimal amounts and leading-zero identifiers stay as text."
        )
    elif extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ValueError("Encrypted uploads are not supported.")
        if len(reader.pages) > 100:
            raise ValueError("PDF uploads support at most 100 pages.")
        for i, page in enumerate(reader.pages):
            stream = page.get_contents()
            if stream is not None and len(stream.get_data()) > 5 * 1024 * 1024:
                raise ValueError("A PDF page exceeds the extraction limit.")
            parts.append(f"[Page {i + 1}]\n" + (page.extract_text(extraction_mode="layout") or ""))
        if not any(re.search(r"[A-Za-zÀ-ž]{3}", s.split("\n", 1)[-1]) for s in parts):
            raise ValueError("This PDF has no readable text. Upload a text-based PDF or Excel file.")
    elif extension == ".docx":
        from docx import Document

        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text]
        for i, t in enumerate(doc.tables):
            item = table(f"t{i + 1}", f"Table {i + 1}", [[c.text for c in row.cells] for row in t.rows])
            if item:
                tables.append(item)
    else:
        parts = [Path(path).read_text(encoding="utf-8-sig")]
    text = "\n".join(parts)
    # Excel tables are retained completely; text outside tables is a bounded
    # preview, explicitly marked rather than misrepresented as the full file.
    truncated = len(text) > MAX_TEXT
    if extension not in (".xlsx", ".csv") and truncated:
        raise ValueError("The uploaded document exceeds the text extraction limit.")
    return {
        "tables": tables,
        "text": text[:MAX_TEXT],
        "text_truncated": truncated,
        "warnings": warnings,
        "format": layout,
    }


def main():
    import resource

    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (40, 40))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_PARSED_BYTES, MAX_PARSED_BYTES))
    try:
        if len(sys.argv)>4 and sys.argv[4]=='--format-only':
            from openpyxl import load_workbook
            from .reference_layout import workbook_layout
            archive_size = guard_zip(sys.argv[1])
            if _xlsx_requires_streaming(sys.argv[1], archive_size):
                result = {"format": extract_xlsx_streaming(sys.argv[1])["format"]}
                Path(sys.argv[3]).write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
                return
            book = load_workbook(sys.argv[1],data_only=True,keep_links=False)
            try:
                if len(book.worksheets)>20 or sum(s.max_row*s.max_column for s in book)>MAX_CELLS or any(s.max_row>MAX_ROWS+50 or s.max_column>MAX_COLS for s in book):
                    raise ValueError("The reference exceeds parsing limits.")
                result = {"format":workbook_layout(book)}
            finally:
                book.close()
        else:
            result = extract(sys.argv[1], sys.argv[2])
        output = Path(sys.argv[3])
        with output.open("w", encoding="utf-8") as stream:
            json.dump(result, stream, ensure_ascii=False, allow_nan=False, separators=(",", ":"))
        if output.stat().st_size > MAX_PARSED_BYTES:
            output.unlink(missing_ok=True)
            raise ValueError("The extracted workbook exceeds 96 MB. Split it into smaller files and upload again.")
    except ValueError as exc:
        Path(sys.argv[3]).write_text(json.dumps({"error": str(exc)}))
        sys.exit(1)
    except Exception:
        Path(sys.argv[3]).write_text(
            json.dumps({"error": "The uploaded file could not be read. Check its format and contents."})
        )
        sys.exit(1)


if __name__ == "__main__":
    main()
