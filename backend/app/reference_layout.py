"""Bounded, presentation-only Excel references; never retain historical figures."""

import re
from xml.etree import ElementTree as ET


def workbook_layout(workbook):
    from openpyxl.styles.colors import COLOR_INDEX

    theme = []
    if workbook.loaded_theme:
        root = ET.fromstring(workbook.loaded_theme)
        scheme = root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
        if scheme is not None:
            for entry in scheme:
                color = next(iter(entry), None)
                theme.append(color.get('lastClr', color.get('val', '000000')) if color is not None else '000000')

    def color(value):
        if value is None:
            return None
        raw = value.rgb if value.type == 'rgb' else (
            theme[value.theme] if value.type == 'theme' and value.theme < len(theme) else
            COLOR_INDEX[value.indexed] if value.type == 'indexed' and value.indexed < len(COLOR_INDEX) else None)
        if not isinstance(raw, str) or not re.fullmatch(r'[0-9A-Fa-f]{6,8}', raw):
            return None
        channels = [int(raw[-6:][i:i+2], 16) for i in (0, 2, 4)]
        tint = value.tint or 0
        channels = [round(c * (1+tint) if tint < 0 else c + (255-c)*tint) for c in channels]
        return '#' + ''.join(f'{min(255,max(0,c)):02X}' for c in channels)

    styles, ids = [], {}

    def style(cell):
        key = cell.style_id
        if key in ids:
            return ids[key]
        if len(styles) >= 256:
            return 0
        alignment = cell.alignment
        item = {'font_name': str(cell.font.name or 'Calibri')[:80],
                'font_size': min(30, max(8, cell.font.sz or 11)),
                'bold': bool(cell.font.bold), 'italic': bool(cell.font.italic),
                'text_wrap': bool(alignment.wrap_text), 'valign': 'vcenter',
                'num_format': str(cell.number_format or 'General')[:200]}
        if alignment.horizontal in ('left', 'center', 'right'):
            item['align'] = alignment.horizontal
        if color(cell.font.color):
            item['font_color'] = color(cell.font.color)
        if cell.fill.patternType == 'solid' and color(cell.fill.fgColor):
            item['bg_color'] = color(cell.fill.fgColor)
        border_styles = {'hair': 7, 'thin': 1, 'medium': 2, 'thick': 5, 'double': 6, 'dashed': 3, 'dotted': 4}
        for edge in ('top', 'bottom', 'left', 'right'):
            side = getattr(cell.border, edge)
            if side and side.style in border_styles:
                item[edge] = border_styles[side.style]
                if color(side.color):
                    item[edge+'_color'] = color(side.color)
        ids[key] = len(styles)
        styles.append(item)
        return ids[key]

    profiles = []
    for sheet in workbook.worksheets[:20]:
        header = next((r for r in range(1, min(30, sheet.max_row)+1)
                       if sum(isinstance(c.value, str) and bool(c.value.strip())
                              for c in sheet[r]) >= 2), None)
        if header is None:
            continue
        # A real Excel table gives the exact header even when a title has several labels.
        if sheet.tables:
            from openpyxl.utils.cell import range_boundaries
            first, header, last, _ = range_boundaries(next(iter(sheet.tables.values())).ref)
        else:
            first = min(c.column for c in sheet[header] if c.value is not None)
            last = max(c.column for row in sheet.iter_rows(min_row=header, max_row=min(header+6,sheet.max_row))
                       for c in row if c.value is not None)
        end = header
        for merge in sheet.merged_cells.ranges:
            if merge.min_row <= header <= merge.max_row and merge.max_row <= header+5:
                end = max(end, merge.max_row)
                last = max(last, merge.max_col)
        # Grouped headers commonly leave the identifier blank below the top level.
        while end < min(header+5, sheet.max_row):
            row = list(sheet.iter_rows(min_row=end+1, max_row=end+1, min_col=first,max_col=last))[0]
            if row[0].value is not None or sum(isinstance(c.value,str) and bool(c.value.strip()) for c in row) < 2:
                break
            end += 1
        last = min(last, first+99)
        from openpyxl.utils import get_column_letter
        body_rows = []
        candidates = list(range(end+1,min(sheet.max_row,end+30)+1)) + list(range(max(end+31,sheet.max_row-19),sheet.max_row+1))
        for r in candidates:
            cells = list(sheet.iter_rows(min_row=r,max_row=r,min_col=first,max_col=last))[0]
            if not isinstance(cells[0].value,str) or not cells[0].value.strip():
                continue
            body_rows.append({'label': cells[0].value[:200], 'styles':[style(c) for c in cells],
                              'height': min(160,max(18,sheet.row_dimensions[r].height or 22))})
        base_row = min(end+1,sheet.max_row)
        frozen_rows, frozen_columns = 0, 0
        if sheet.freeze_panes:
            from openpyxl.utils.cell import coordinate_to_tuple
            frozen_rows, frozen_columns = coordinate_to_tuple(sheet.freeze_panes)
        title_cell = next((c for row in sheet.iter_rows(min_row=1,max_row=max(1,header-1)) for c in row if isinstance(c.value,str) and c.value.strip()), sheet.cell(header,first))
        profiles.append({'name': sheet.title, 'title_style':style(title_cell), 'start_column':min(first-1,3),
                         'headers': [[str(sheet.cell(r,c).value)[:300] if sheet.cell(r,c).value is not None else None
                                      for c in range(first,last+1)] for r in range(header,end+1)],
                         'header_styles':[[style(sheet.cell(r,c)) for c in range(first,last+1)] for r in range(header,end+1)],
                         'header_heights':[min(100,max(20,sheet.row_dimensions[r].height or 26)) for r in range(header,end+1)],
                         'merges': [[m.min_row-header,m.min_col-first,m.max_row-header,m.max_col-first]
                                    for m in sheet.merged_cells.ranges
                                    if header<=m.min_row<=m.max_row<=end and first<=m.min_col<=m.max_col<=last],
                         'column_widths':[min(60,max(8,sheet.column_dimensions[get_column_letter(c)].width or 14)) for c in range(first,last+1)],
                         'body_styles':[style(sheet.cell(base_row,c)) for c in range(first,last+1)],
                         'body_rows':body_rows,'landscape':sheet.page_setup.orientation=='landscape',
                         'freeze_columns':max(0,min(last-first+1,frozen_columns-first)),
                         'freeze_headers':bool(frozen_rows), 'zoom':sheet.sheet_view.zoomScale or 85})
    return {'layout_version':2,'sheet_layouts':profiles,'styles':styles,
            'layout_policy':'Presentation evidence only. Author all headers for the requested period and language; never reuse historical figures or formulas. Body labels select appearance only, never eligible data.'}
