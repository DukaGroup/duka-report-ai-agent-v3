import io
import asyncio
import json
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from app.connectors import DatabaseUnavailable
from app.database_tools import DatabaseQuery, compile_select
from app.reports import bounded_sample
from conftest import HEADERS
from test_database_chat import DESCRIPTIONS, fixture_source, query_plan, call_tool
from test_excel_reports import subscribe, wait_report
from test_sqlserver import Driver, connector


@pytest.mark.asyncio
async def test_full_sql_cursor_fetches_every_row_without_preview_top_or_offset():
    source_rows = [{"customer": i, "revenue": Decimal("100.25")} for i in range(6856)]
    driver = Driver(source_rows)
    database = connector(driver)

    async def describe(schema, table):
        return DESCRIPTIONS["s"]

    database.describe = describe
    plan = query_plan(select=[{"column": "s.CustomerId", "alias": "customer"}], group_by=["s.CustomerId"], limit=2)
    result = await database.report_query(plan)
    assert len(result["rows"]) == 6856 and not result["truncated"]
    assert result["rows"][-1] == {"customer": 6855, "revenue": "100.25"}
    assert "TOP (" not in driver.executed[0] and "OFFSET" not in driver.executed[0]
    assert "GROUP BY" in driver.executed[0] and driver.closed and driver.rolled_back
    limited_sql, _, _ = compile_select(DatabaseQuery.model_validate(plan), DESCRIPTIONS)
    assert "SELECT TOP (3)" in limited_sql


def test_excel_contains_every_customer_and_uses_full_data_formulas_and_summaries(client, admin):
    subscribe(client)
    source, database = fixture_source(client)
    total_rows = 6856
    full = [{"customer": i, "revenue": "100.25"} for i in range(total_rows)]

    def fetch_all(sql, parameters):
        assert "SELECT TOP" not in sql
        return full if "[customer]" in sql else [{"revenue": str(Decimal("100.25") * total_rows)}]

    database._run_full = fetch_all
    plan = query_plan(
        select=[{"column": "s.CustomerId", "alias": "customer"}],
        group_by=["s.CustomerId"],
        order_by=[{"column": "customer"}],
        limit=2,
    )
    identifier = call_tool(
        client,
        admin,
        "create_database_report",
        {
            "source_id": source["id"],
            "query": plan,
            "title": "Full sales",
            "analysis": {
                "name": "Every customer",
                "kind": "breakdown",
                "dimension": "customer",
                "measures": ["revenue"],
            },
            "analyses": [
                {
                    "name": "Second full breakdown",
                    "query": plan,
                    "kind": "breakdown",
                    "dimension": "customer",
                    "measures": ["revenue"],
                }
            ],
        },
    )["report_id"]
    report = wait_report(client, identifier)
    assert report["status"] == "completed", report
    snapshot = client.app.state.db.one("SELECT * FROM dataset_snapshots WHERE id=?", (report["snapshot_id"],))
    assert len(json.loads(snapshot["rows_json"])) == total_rows
    provenance = json.loads(snapshot["provenance_json"])
    assert len(provenance["analyses"][0]["rows"]) == total_rows and provenance["analysis_errors"] == []
    response = client.get(f"/api/reports/{identifier}/export/xlsx")
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    tables = [sheet for sheet in workbook if sheet.tables]
    assert len(tables) == 2
    for sheet in tables:
        table = next(iter(sheet.tables.values()))
        first_col, header_row, last_col, last_row = range_boundaries(table.ref)
        assert first_col == 1 and last_col == 3 and last_row - header_row == total_rows
        assert [sheet.cell(i, 1).value for i in range(header_row + 1, last_row + 1)] == list(range(total_rows))
        assert sheet.cell(last_row, 2).value == 100.25
    overview = workbook.worksheets[0]
    assert any("SUM(" in str(cell.value) and str(last_row) in str(cell.value) for row in overview for cell in row)
    assert len(bounded_sample(full)) <= 100


@pytest.mark.asyncio
async def test_export_capacity_stops_instead_of_returning_partial_rows():
    driver = Driver([{"revenue": Decimal("10")}] * 1001)
    database = connector(driver)
    database.settings.report_max_rows = 1000
    with pytest.raises(DatabaseUnavailable, match="no partial report"):
        await asyncio.to_thread(database._run_full, "SELECT revenue FROM dbo.Sales", {})
    assert driver.closed and driver.rolled_back


@pytest.mark.asyncio
async def test_compact_utf8_result_size_matches_snapshot_and_never_changes_values():
    rows = [{"Përshkrimi": "Shitje në Shqipëri", "revenue": "100.25"} for _ in range(1200)]
    compact = json.dumps(rows, ensure_ascii=False, separators=(",", ":")).encode()
    driver = Driver(rows)
    database = connector(driver)
    database.settings.report_max_bytes = len(compact) + 1
    result = await asyncio.to_thread(database._run_full, "SELECT description,revenue FROM dbo.Sales", {})
    assert result == rows
    assert len(json.dumps(result, ensure_ascii=False, separators=(",", ":")).encode()) == len(compact)


def test_manual_full_report_has_bounded_model_context_and_full_excel(client, admin, monkeypatch):
    from app.provider import CodexRuntime

    subscribe(client)
    rows = [{"revenue": i, "cost": 1} for i in range(2500)]
    original = CodexRuntime.stream
    prompts = []

    async def capture(self, thread_id, content, model=None, output_schema=None):
        if output_schema:
            prompts.append(json.loads(content))
        async for event in original(self, thread_id, content, model, output_schema):
            yield event

    monkeypatch.setattr(CodexRuntime, "stream", capture)
    result = client.post("/api/reports", headers=HEADERS, json={"title": "All rows", "rows": rows}).json()
    report = wait_report(client, result["id"])
    assert report["status"] == "completed"
    context = prompts[0]
    assert len(context["data"]) < 2500
    assert context["data_metadata"]["row_count"] == 2500
    assert "computed_summary" not in context
    assert "numeric_column_sums" not in context["data_metadata"]
    artifact = load_workbook(io.BytesIO(client.get(f"/api/reports/{result['id']}/export/xlsx").content), data_only=True)
    data = next(sheet for sheet in artifact if sheet.tables)
    _, header, _, end = range_boundaries(next(iter(data.tables.values())).ref)
    assert end - header == 2500 and data.cell(end, 1).value == 2499
