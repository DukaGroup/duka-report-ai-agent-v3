import io
import json

import pytest
from openpyxl import load_workbook

from app.balance_reconciliation import BalanceRequest, bridge_rows, columns_for
from app.balance_workbook import balance_contract
from app.workbooks import ReportDataError
from app.workbook_validation import validate_workbook
from app.workbook_runner import WorkbookRunner
from app.config import Settings
from test_excel_reports import workbook_fixture
from test_database_chat import fixture_source
from conftest import HEADERS


OPTIONS = {"company": "DTR", "as_of_date": "2026-10-09"}


def line(**changes):
    return {
        "party": "supplier",
        "partner_code": "V1",
        "partner_name": "Test supplier",
        "currency": "ALL",
        "client": "400",
        "company": "DTR",
        "year": "2026",
        "document": "100",
        "line": "001",
        "posting_date": "20261001",
        "entry_date": "20261002",
        "clearing_date": "00000000",
        "clearing_entry_date": None,
        "sign": "H",
        "amount_local": "100.00",
        **changes,
    }


def example():
    return [
        line(
            document="payment",
            sign="S",
            amount_local="209368.25",
            clearing_date="20261001",
            clearing_document="reversal",
            clearing_entry_date="20261014",
            clearing_transaction="FB08",
        ),
        line(
            document="reversal",
            amount_local="209368.25",
            entry_date="20261014",
            clearing_date="20261001",
            clearing_document="reversal",
            clearing_entry_date="20261014",
            transaction="FB08",
        ),
        line(document="invoice", amount_local="11029.23"),
        line(
            party="customer",
            partner_code="C1",
            partner_name="Customer",
            document="sale",
            sign="S",
            amount_local="1000001.00",
        ),
        line(
            party="customer",
            partner_code="C1",
            partner_name="Customer",
            document="receipt",
            sign="H",
            amount_local="1000000.00",
            entry_date="20261014",
        ),
    ]


def test_reconstructs_backdated_payment_reversal_without_double_counting_closed_credit():
    rows, changes = bridge_rows(example(), OPTIONS)
    customer, supplier = rows
    assert customer["original_balance"] == "1000001.00" and customer["late_postings"] == "-1000000.00"
    assert customer["restated_balance"] == "1.00"
    assert supplier["original_balance"] == "-198339.02"
    assert supplier["backdated_reversals"] == "209368.25"
    assert supplier["restated_balance"] == "11029.23"
    assert all(r["residual"] == "0.00" for r in rows)
    assert len(changes) == 3
    assert next(r for r in changes if r["document"] == "reversal")["change"] == "0"


def test_regular_backdated_clearing_is_separate_and_baseline_cutoff_is_inclusive():
    rows, _ = bridge_rows([line(clearing_date="20261008", clearing_entry_date="20261010")], OPTIONS)
    assert rows[0]["original_balance"] == "100.00" and rows[0]["backdated_clearing"] == "-100.00"
    rows, _ = bridge_rows([line(entry_date="20261009")], OPTIONS)
    assert rows[0]["original_balance"] == "100.00" and rows[0]["late_postings"] == "0"
    rows, _ = bridge_rows(example(), {**OPTIONS, "knowledge_date": "2026-10-14"})
    assert all(r["original_balance"] == r["restated_balance"] for r in rows)


def test_missing_history_duplicates_and_invalid_dates_fail_instead_of_inventing_original_balance():
    for evidence in (
        [line(entry_date=None)],
        [line(clearing_date="20261001")],
        [line(), line()],
        [line(posting_date="invalid")],
    ):
        with pytest.raises(ReportDataError):
            bridge_rows(evidence, OPTIONS)
    with pytest.raises(ValueError):
        BalanceRequest.model_validate({**OPTIONS, "knowledge_date": "2026-10-08"})


def snapshot_fixture(language="en"):
    evidence = example()
    rows, changes = bridge_rows(evidence, OPTIONS)
    report, snapshot = workbook_fixture(rows)
    provenance = {
        "columns": columns_for(rows),
        "analysis": {"name": "Balance bridge"},
        "reconciliation": {"method": "sap_balance_bridge_v1", "request": OPTIONS, "currency": "ALL"},
        "analyses": [
            {"rows": changes, "columns": columns_for(changes), "options": {"name": "Balance adjustments"}},
            {"rows": evidence, "columns": columns_for(evidence), "options": {"name": "SAP balance evidence"}},
        ],
    }
    snapshot["provenance_json"] = json.dumps(provenance)
    document = json.loads(report["content_json"])
    document["workbook"] = {
        "language": language,
        "labels": [],
        "analyses": [
            {
                "index": i,
                "name": name,
                "scope": "Fixture scope",
                "query_explanation": "Fixture evidence",
                "columns": [{"column": c["name"], "label": c["name"]} for c in columns_for(data)],
            }
            for i, name, data in [(0, "Bridge", rows), (1, "Adjustments", changes), (2, "Evidence", evidence)]
        ],
        "python_script": "def build(report):\n    s = report.sheet('Review')\n    s.title()\n    s.narrative()",
    }
    report["content_json"] = json.dumps(document)
    return report, snapshot, document


@pytest.mark.parametrize("language,name", [("en", "Balance reconciliation"), ("sq", "Rakordimi i gjendjeve")])
async def test_worker_adds_mandatory_localized_bridge_and_checks_every_adjustment(tmp_path, language, name):
    report, snapshot, document = snapshot_fixture(language)
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    content, verified = await WorkbookRunner(settings).build(report, snapshot)
    assert verified["source_rows"] == [2, 3, 5]
    w = load_workbook(io.BytesIO(content), data_only=True)
    assert w.sheetnames[0] == name
    assert w[name]["E10"].value == -198339.02 and w[name]["G10"].value == 209368.25
    assert w[name]["I10"].value == 11029.23 and w[name]["J10"].value == 0
    assert len(w[name]._charts) == 1
    w[name]["G10"] = 0
    out = io.BytesIO()
    w.save(out)
    with pytest.raises(ReportDataError, match="balance totals"):
        validate_workbook(out.getvalue(), document, snapshot, [], True)


def test_source_bridge_must_match_immutable_document_evidence():
    _, snapshot, _ = snapshot_fixture()
    rows = json.loads(snapshot["rows_json"])
    rows[0]["late_postings"] = "0"
    snapshot["rows_json"] = json.dumps(rows)
    with pytest.raises(ReportDataError, match="document evidence"):
        balance_contract(snapshot)




def test_new_reports_reject_fixed_balance_calculator_before_querying(client, admin, monkeypatch):
    source, connector = fixture_source(client)
    async def forbidden(**kwargs):
        raise AssertionError("Retired report requests cannot query a database")
    monkeypatch.setattr(connector, "balance_reconciliation", forbidden)
    response = client.post("/api/reports", headers=HEADERS, json={
        "source_id": source["id"], "database": "Fixture", "title": "Old fixed calculation",
        "balance_reconciliation": OPTIONS,
    })
    assert response.status_code == 422 and "retired" in response.text
    assert connector.calls == []
    assert client.app.state.db.all("SELECT * FROM reports") == []
    assert client.app.state.db.all("SELECT * FROM dataset_snapshots") == []
