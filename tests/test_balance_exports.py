"""The same requested-scope contract governs every balance export/presentation."""

import copy
import io
import json
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.balance_exports import CUTOFF_FIELDS, DOCUMENT_FIELDS, prepare_balance_export
from app.balance_reconciliation import bridge_rows, columns_for
from app.balance_workbook import balance_contract
from app.config import Settings
from app.report_exports import public_snapshot
from app.reports import csv_snapshot
from app.workbook_runner import WorkbookRunner
from app.workbook_validation import validate_workbook
from app.workbooks import ReportDataError, excel_report
from conftest import HEADERS
from test_balance_calculations import BalanceDatabase, registered_database
from test_balance_reconciliation import OPTIONS, line
from test_excel_reports import workbook_fixture


def raw_result(evidence, options):
    rows, changes = bridge_rows(evidence, options)
    return {
        "rows": rows, "changes": changes, "evidence": evidence, "columns": columns_for(rows),
        "reconciliation": {"method": "sap_balance_bridge_v1", "request": options,
                           "currency": evidence[0]["currency"] if evidence else "ALL", "limitations": []},
        "query": options, "scope": "Requested scope", "truncated": False,
    }


def prepared_snapshot(evidence, options):
    view, audit = prepare_balance_export(raw_result(evidence, options), options)
    report, snapshot = workbook_fixture(view["rows"])
    cutoff = view["reconciliation"]["export_mode"] == "cutoff"
    supporting = [("Eligible balance documents", view["evidence"], DOCUMENT_FIELDS)] if cutoff else [
        ("Balance adjustments", view["changes"], None), ("SAP balance evidence", view["evidence"], None),
    ]
    provenance = {
        "columns": view["columns"], "query": view["query"], "scope": view["scope"],
        "reconciliation": view["reconciliation"],
        "analyses": [{"rows": rows, "columns": columns_for(rows, fields), "options": {"name": name}}
                     for name, rows, fields in supporting],
    }
    if audit is not None:
        provenance["_internal"] = {"balance_audit": audit}
    snapshot["provenance_json"] = json.dumps(provenance)
    return report, snapshot


def evidence_fixture(tmp_path):
    return [r for r in BalanceDatabase(Settings(data_dir=tmp_path, env_file=None), None).evidence if r["client"] == "400"]


@pytest.mark.parametrize("asof,company,client,currency,loan_account", [
    ("2026-08-25", "ECO", "400", "ALL", "4610041000"),
    ("2026-09-02", "ABCD", "500", "EUR", "9000100001"),
    ("2026-10-09", "XYZ", "700", "USD", "1234567890"),
])
def test_cutoff_projection_is_universal_and_retains_correct_backdated_clearing_membership(
    tmp_path, asof, company, client, currency, loan_account,
):
    evidence = evidence_fixture(tmp_path)
    shift = date.fromisoformat(asof) - date(2026, 10, 9)
    for row in evidence:
        row.update(company=company, client=client, currency=currency)
        if row["account"] == "4610041000":
            row["account"] = loan_account
        for field in ("posting_date", "entry_date", "clearing_date", "clearing_entry_date"):
            value = row.get(field)
            if value and value != "00000000":
                row[field] = (date.fromisoformat(value[:4] + "-" + value[4:6] + "-" + value[6:8]) + shift).strftime("%Y%m%d")
    options = {"company": company, "client": client, "as_of_date": asof, "exclude_supplier_accounts": [loan_account]}
    view, audit = prepare_balance_export(raw_result(evidence, options), options)
    assert view["reconciliation"]["export_mode"] == "cutoff"
    assert set(r["document"] for r in view["evidence"]) == {"invoice", "cleared", "mixed-trade", "sale"}
    assert all(set(r) == set(DOCUMENT_FIELDS) for r in view["evidence"])
    assert all(set(r) == set(CUTOFF_FIELDS) for r in view["rows"])
    assert view["reconciliation"]["totals"] == {
        "supplier/" + currency: {"original_balance": "185.00"},
        "customer/" + currency: {"original_balance": "200"},
    }
    assert view["changes"] == []
    assert any(r["entry_date"] > asof.replace("-", "") for r in audit["evidence"])
    assert all(r["account"] != loan_account for r in view["evidence"])
    assert "clearing_entry_date" not in DOCUMENT_FIELDS


@pytest.mark.parametrize("custom", [False, True])
@pytest.mark.parametrize("language,name", [("en", "Cutoff balances"), ("sq", "Gjendjet në datën kufi")])
async def test_both_workbook_paths_show_only_cutoff_balances_and_eligible_documents(tmp_path, custom, language, name):
    options = {**OPTIONS, "exclude_supplier_accounts": ["4610041000"]}
    report, snapshot = prepared_snapshot(evidence_fixture(tmp_path), options)
    document = json.loads(report["content_json"])
    document["workbook"] = {
        "language": language, "labels": [], "analyses": [],
        "python_script": "def build(report):\n    s=report.sheet('Review')\n    s.table(0)\n    s.table(1)\n",
    }
    provenance = json.loads(snapshot["provenance_json"])
    document["workbook"]["analyses"] = [
        {"index": index, "name": "Requested data " + str(index), "scope": "Requested scope",
         "query_explanation": "Cutoff and verified exclusions",
         "columns": [{"column": column["name"], "label": column["name"]} for column in dataset["columns"]]}
        for index, dataset in enumerate([provenance, *provenance["analyses"]])
    ]
    report["content_json"] = json.dumps(document)
    if custom:
        settings = Settings(data_dir=tmp_path, env_file=None)
        settings.prepare()
        content, verified = await WorkbookRunner(settings).build(report, snapshot)
        assert verified["source_rows"] == [3, 4]
    else:
        document["workbook"]["python_script"] = ""
        report["content_json"] = json.dumps(document)
        content = excel_report(report, snapshot)
        validate_workbook(content, document, snapshot, [], False)
    w = load_workbook(io.BytesIO(content), data_only=True)
    assert w.sheetnames[0] == name
    assert w[name]["E10"].value == 185 and w[name]["E11"].value == 200
    assert "Balance reconciliation" not in w.sheetnames and "Rakordimi i gjendjeve" not in w.sheetnames
    assert "Balance adjustments" not in w.sheetnames and "Ndryshimet e gjendjeve" not in w.sheetnames
    values = [cell.value for sheet in w for row in sheet for cell in row]
    assert not any(value in ("late", "receipt", "loan", "20261010") for value in values if isinstance(value, str))
    assert "restated_balance" not in values and "late_postings" not in values
    assert "clearing_entry_date" not in values
    # The arithmetic control is validated independently of the model-authored views.
    w[name]["E10"] = 999
    output = io.BytesIO()
    w.save(output)
    with pytest.raises(ReportDataError, match="cutoff balance totals"):
        validate_workbook(output.getvalue(), document, snapshot, [], custom)


def test_explicit_reconciliation_preserves_the_full_comparison(tmp_path):
    report, snapshot = prepared_snapshot(evidence_fixture(tmp_path), {
        **OPTIONS, "exclude_supplier_accounts": ["4610041000"], "include_reconciliation": True,
    })
    provenance = json.loads(snapshot["provenance_json"])
    assert provenance["reconciliation"]["export_mode"] == "reconciliation"
    assert "_internal" not in provenance
    assert "restated_balance" in json.loads(snapshot["rows_json"])[0]
    w = load_workbook(io.BytesIO(excel_report(report, snapshot)), data_only=True)
    assert "Balance reconciliation" in w.sheetnames and "Balance adjustments" in w.sheetnames
    assert w["Balance reconciliation"]["E10"].value == 185
    assert w["Balance reconciliation"]["I10"].value == 90


@pytest.mark.parametrize("evidence", [[], [
    line(document="invoice", account="4011", amount_local="10", entry_time="235959"),
    line(document="payment", account="4011", amount_local="10", sign="S"),
    line(document="later-only", account="4011", partner_code="LATE", entry_date="20261010"),
]])
def test_empty_and_net_zero_balances_keep_only_the_eligible_membership(evidence):
    report, snapshot = prepared_snapshot(evidence, OPTIONS)
    meta, rows, _ = balance_contract(snapshot)
    assert all(Decimal(total["original_balance"]) == 0 for total in meta["totals"].values())
    assert set(meta["totals"]) == {"supplier/ALL", "customer/ALL"}
    assert len(rows) == (1 if evidence else 0)
    assert meta["document_count"] == (2 if evidence else 0)
    w = load_workbook(io.BytesIO(excel_report(report, snapshot)), data_only=True)
    assert w["Cutoff balances"]["E10"].value == 0
    assert w["Cutoff balances"]["E11"].value == 0


def test_generic_public_snapshot_boundary_strips_private_audit_without_changing_business_data():
    snapshot = {"rows_json": '[{"sales":"125.00"}]', "provenance_json": json.dumps({
        "source": "Any database", "columns": [{"name": "sales", "type": "decimal"}],
        "_internal": {"audit": "Private history"},
    })}
    public = public_snapshot(snapshot)
    assert public["rows_json"] == snapshot["rows_json"]
    assert json.loads(public["provenance_json"]) == {
        "source": "Any database", "columns": [{"name": "sales", "type": "decimal"}],
    }


def test_private_history_never_enters_public_metadata_and_csv_contains_no_revised_fields(tmp_path):
    report, snapshot = prepared_snapshot(evidence_fixture(tmp_path), {**OPTIONS, "exclude_supplier_accounts": ["4610041000"]})
    before = copy.deepcopy(snapshot)
    public = public_snapshot(snapshot)
    assert "_internal" not in json.loads(public["provenance_json"])
    assert "_internal" in json.loads(snapshot["provenance_json"])
    assert snapshot == before  # The reconstruction snapshot is retained intact.
    assert "late_postings" not in csv_snapshot(snapshot) and "restated_balance" not in csv_snapshot(snapshot)
    assert "original_balance" in csv_snapshot(snapshot)


def test_scope_validation_rejects_injected_dates_accounts_missing_documents_and_false_signed_amounts(tmp_path):
    report, snapshot = prepared_snapshot(evidence_fixture(tmp_path), {**OPTIONS, "exclude_supplier_accounts": ["4610041000"]})
    for mutation in ("date", "account", "missing", "signed"):
        p = json.loads(snapshot["provenance_json"])
        docs = p["analyses"][0]["rows"]
        if mutation == "date":
            docs[0]["entry_date"] = "20261010"
        elif mutation == "account":
            docs[0]["account"] = "4610041000"
        elif mutation == "missing":
            docs.pop()
        else:
            docs[0]["original_balance"] = "999"
        corrupted = {**snapshot, "provenance_json": json.dumps(p)}
        with pytest.raises(ReportDataError, match="scope|reconcile"):
            balance_contract(corrupted)
        with pytest.raises(ReportDataError, match="scope|reconcile"):
            excel_report(report, corrupted)


async def test_generated_script_cannot_access_later_columns(tmp_path):
    report, snapshot = prepared_snapshot(evidence_fixture(tmp_path), OPTIONS)
    document = json.loads(report["content_json"])
    document["workbook"] = {"language": "en", "labels": [], "analyses": [],
                            "python_script": "def build(report):\n    s=report.sheet('Bad')\n    s.table(0,columns=['restated_balance'])"}
    report["content_json"] = json.dumps(document)
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    with pytest.raises(ReportDataError, match="source column"):
        await WorkbookRunner(settings).build(report, snapshot)


def test_old_export_parameters_cannot_reenable_a_fixed_calculator(client, admin):
    source, connector = registered_database(client)
    for include in (False, True):
        response = client.post("/api/reports", headers=HEADERS, json={
            "title": "Old balances", "source_id": source["id"], "database": "Fixture",
            "balance_reconciliation": {**OPTIONS, "include_reconciliation": include},
        })
        assert response.status_code == 422 and "retired" in response.text
    assert connector.calls == []
    assert client.app.state.db.all("SELECT * FROM reports") == []
