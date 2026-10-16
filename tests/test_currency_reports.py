import io
import json
import zipfile

from openpyxl import load_workbook

from app.currencies import currency_rows
from app.provider import CodexRuntime
from app.workbooks import current_workbook, excel_report
from conftest import HEADERS
from test_excel_reports import subscribe, wait_report, workbook_fixture
from test_prompt_guidance import last_input


def fixture():
    report, snapshot = workbook_fixture(
        [
            {"Currency Code": "", "revenue": 10, "note": ""},
            {"Currency Code": None, "revenue": 20, "note": None},
            {"Currency Code": " \t", "revenue": 30, "note": ""},
            {"Currency Code": "EUR", "revenue": 40, "note": ""},
            {"Currency Code": "USD", "revenue": 50, "note": ""},
        ],
        {"name": "Currencies", "kind": "data", "measures": []},
    )
    provenance = json.loads(snapshot["provenance_json"])
    provenance["columns"] = [
        {"name": "Currency Code", "type": "nvarchar"},
        {"name": "revenue", "type": "decimal"},
        {"name": "note", "type": "nvarchar"},
    ]
    snapshot["provenance_json"] = json.dumps(provenance)
    return report, snapshot


def test_excel_displays_all_for_blank_currency_codes_without_changing_source_values():
    report, snapshot = fixture()
    original = dict(snapshot)
    workbook = load_workbook(io.BytesIO(excel_report(report, snapshot)), data_only=True)
    sheet = workbook.worksheets[1]
    assert [sheet.cell(row, 1).value for row in range(6, 11)] == ["ALL", "ALL", "ALL", "EUR", "USD"]
    assert [sheet.cell(row, 2).value for row in range(6, 11)] == [10, 20, 30, 40, 50]
    assert all(sheet.cell(row, 3).value in (None, "") for row in range(6, 11))
    assert any(cell.value == "Workspace base currency: ALL" for row in workbook.worksheets[0] for cell in row)
    assert any("Blank currency codes" in str(cell.value) for row in sheet for cell in row)
    assert snapshot == original


def test_currency_aliases_and_albanian_column_names_are_recognized():
    rows = [{"ccy": "", "Monedha": None, "Kodi i monedhës": "  ", "Currency Exchange Rate": None}]
    query = {"select": [{"column": "s.Currency Code", "alias": "ccy"}]}
    assert currency_rows(rows, query=query) == [
        {"ccy": "ALL", "Monedha": "ALL", "Kodi i monedhës": "ALL", "Currency Exchange Rate": None}
    ]
    report, snapshot = workbook_fixture(rows, {"name": "Currency aliases"})
    provenance = {"columns": [], "query": query, "analysis": {"name": "Currency aliases"}}
    snapshot["provenance_json"] = json.dumps(provenance)
    workbook = load_workbook(io.BytesIO(excel_report(report, snapshot)), data_only=True)
    assert [workbook.worksheets[1].cell(6, col).value for col in range(1, 5)] == ["ALL", "ALL", "ALL", None]


def test_background_model_receives_all_labels_and_original_snapshot_is_retained(client, admin):
    subscribe(client)
    rows = [{"Currency Code": "", "revenue": 12}, {"Currency Code": "EUR", "revenue": 3}]
    report = client.post(
        "/api/reports", headers=HEADERS,
        json={"title": "Currency labels", "instructions": "Më jep raportin e shitjeve.", "rows": rows},
    ).json()
    completed = wait_report(client, report["id"])
    assert completed["status"] == "completed", completed
    prompt = json.loads(last_input(client, admin["id"])[0]["text"])
    assert prompt["data"] == [{"Currency Code": "ALL", "revenue": 12}, {"Currency Code": "EUR", "revenue": 3}]
    assert prompt["currency_columns"] == ["Currency Code"]
    assert "confirmed workspace base currency" in prompt["currency_instructions"]
    snapshot = client.app.state.db.one("SELECT * FROM dataset_snapshots WHERE id=?", (report["snapshot_id"],))
    assert json.loads(snapshot["rows_json"]) == rows
    content = client.get(f"/api/reports/{report['id']}/export/xlsx").content
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    assert workbook.worksheets[1]["A6"].value == "ALL"
    assert any(
        cell.value == "Monedha bazë e hapësirës: ALL" for row in workbook.worksheets[0] for cell in row
    )


def test_existing_export_is_refreshed_once_without_model_or_database_queries(client, admin, monkeypatch):
    subscribe(client)
    report = client.post(
        "/api/reports", headers=HEADERS,
        json={"title": "Existing report", "rows": [{"Currency Code": "", "revenue": 10}]},
    ).json()
    assert wait_report(client, report["id"])["status"] == "completed"
    path = client.app.state.settings.data_dir / "reports" / (report["id"] + ".xlsx")
    assert current_workbook(path)
    # Simulate the old presentation cache without its current version marker.
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as archive, zipfile.ZipFile(buffer, "w") as old:
        for name in archive.namelist():
            if name != "docProps/custom.xml":
                old.writestr(name, archive.read(name))
    path.write_bytes(buffer.getvalue())
    assert not current_workbook(path)

    def unexpected(*args, **kwargs):
        raise AssertionError("Presentation refresh must not call the model or repeat generation")

    monkeypatch.setattr(CodexRuntime, "stream", unexpected)
    url = f"/api/reports/{report['id']}/export/xlsx"
    response = client.get(url)
    assert response.status_code == 200
    assert current_workbook(path)
    assert load_workbook(io.BytesIO(response.content), data_only=True).worksheets[1]["A6"].value == "ALL"
    monkeypatch.setattr("app.main.rebuild_excel", unexpected)
    assert client.get(url).content == response.content
