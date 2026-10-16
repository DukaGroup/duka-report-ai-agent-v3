import asyncio
import io
import json
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.connectors import ConnectorRegistry, SourceNotConfigured
from app.db import Database
from app.provider import BASE_INSTRUCTIONS, REPORT_TOOLS
from app.source_secrets import SourceSecrets
from conftest import HEADERS
from test_application import create_member
from test_database_chat import FixtureDatabase, call_tool, fixture_source, query_plan
from test_excel_reports import subscribe, wait_report
from test_sqlserver import Driver, connector


async def test_database_catalog_is_paged_accessible_online_business_databases_only():
    driver = Driver([{"name": "Archive", "is_default": 0}, {"name": "Live", "is_default": 1}])
    result = await connector(driver).databases(limit=1)
    assert result == {"databases": [{"name": "Archive", "is_default": 0}], "next_offset": 1}
    sql, params = driver.executed
    assert "HAS_DBACCESS(name)=1" in sql and "database_id>4" in sql and "state=0" in sql
    assert params == {"offset": 0, "size": 2}
    assert driver.options["database"] == "master"


async def test_database_switch_is_bound_authorized_and_does_not_mutate_concurrent_targets():
    driver = Driver([{"name": "Archive] ; --"}])
    original = connector(driver)
    target = await original.in_database("Archive] ; --")
    assert target is not original and target.settings.sqlserver_database == "Archive] ; --"
    assert original.settings.sqlserver_database == "private-database"
    assert "Archive] ; --" not in driver.executed[0]
    assert driver.executed[1] == {"database": "Archive] ; --"}
    driver.rows, driver.position = [], 0
    with pytest.raises(SourceNotConfigured, match="cannot access"):
        await original.in_database("Forbidden")
    for invalid in ("", "x" * 129, "invalid\x00name", 1):
        with pytest.raises(ValueError):
            await original.in_database(invalid)


def test_generic_env_does_not_override_distinct_saved_secondary_source(tmp_path):
    env = tmp_path / ".env"
    env.write_text("SQLSERVER_HOST=default-host\nSQLSERVER_DATABASE=Default\nSQLSERVER_USER=default-user\nSQLSERVER_PASSWORD=default-secret\n")
    settings = Settings(data_dir=tmp_path / "data", env_file=env)
    settings.prepare()
    db = Database(settings.data_dir / "application.sqlite3")
    db.initialize()
    for identifier in ("first", "second"):
        db.execute("INSERT INTO data_sources(id,name,kind,created_at) VALUES(?,?,'sqlserver','now')", (identifier, identifier))
    secrets = SourceSecrets(settings, db)
    values = {"SQLSERVER_HOST": "separate-host", "SQLSERVER_USER": "separate-user", "SQLSERVER_PASSWORD": "separate-secret", "SQLSERVER_DATABASE": "Separate", "SQLSERVER_PORT": "1433"}
    for identifier in ("first", "second"):
        db.execute("INSERT INTO source_credentials VALUES(?,?,'now')", (identifier, secrets.encrypt(identifier, values)))
    registry = ConnectorRegistry(settings, db)
    assert registry.source_settings("first").sqlserver_database == "Default"
    assert registry.source_settings("second").sqlserver_database == "Separate"
    assert registry.public_source(db.one("SELECT * FROM data_sources WHERE id='second'"))["configuration_origin"] == "saved"


def test_inventory_discovers_other_sources_even_if_one_fails_and_enforces_grants(client, admin, monkeypatch):
    first, first_connector = fixture_source(client)
    second, second_connector = fixture_source(client)

    async def failed(**kwargs):
        raise SourceNotConfigured("Database catalog unavailable.")

    async def available(offset=0, limit=200):
        return {"databases": [{"name": "Live" if offset == 0 else "History", "is_default": offset == 0}], "next_offset": 200 if offset == 0 else None}

    monkeypatch.setattr(first_connector, "databases", failed)
    monkeypatch.setattr(second_connector, "databases", available)
    result = call_tool(client, admin, "discover_databases", {})
    assert not result["catalog_complete"]
    assert result["sources"][0]["error"] and not result["sources"][0]["catalog_complete"]
    assert [d["name"] for d in result["sources"][1]["databases"]] == ["Live", "History"]
    member = create_member(client)
    assert call_tool(client, member, "discover_databases", {})["sources"] == []
    assert "error" in call_tool(client, member, "list_databases", {"source_id": second["id"]})
    client.app.state.db.execute("INSERT INTO source_grants VALUES(?,?)", (second["id"], member["id"]))
    assert [s["id"] for s in call_tool(client, member, "discover_databases", {})["sources"]] == [second["id"]]


class MultipleDatabaseFixture(FixtureDatabase):
    async def in_database(self, database):
        if database not in ("Live", "History"):
            raise SourceNotConfigured("Database is unavailable or the reporting login cannot access it.")
        return MultipleDatabaseFixture(replace(self.settings, sqlserver_database=database), self.catalog)

    def _run_full(self, sql, parameters):
        return [{"revenue": "10.25" if self.settings.sqlserver_database == "Live" else "20.50"}]


def test_one_excel_can_include_two_databases_and_another_source_with_provenance(client, admin):
    subscribe(client)
    first, _ = fixture_source(client)
    second, _ = fixture_source(client)
    state = client.app.state
    for source in (first, second):
        state.connectors.register(source["id"], MultipleDatabaseFixture(replace(state.settings, sqlserver_database="Live"), state.connectors.catalog))
    result = client.post("/api/reports", headers=HEADERS, json={
        "title": "Cross-database report", "source_id": first["id"], "database": "Live", "query": query_plan(),
        "analyses": [
            {"name": "Historical sales", "database": "History", "query": query_plan()},
            {"name": "Other source", "source_id": second["id"], "query": query_plan()},
        ],
    })
    assert result.status_code == 201, result.text
    report = wait_report(client, result.json()["id"])
    assert report["status"] == "completed", report
    snapshot = state.db.one("SELECT * FROM dataset_snapshots WHERE id=?", (report["snapshot_id"],))
    provenance = json.loads(snapshot["provenance_json"])
    assert provenance["database"] == "Live"
    assert [(a["source_id"], a["database"], a["rows"][0]["revenue"]) for a in provenance["analyses"]] == [(first["id"], "History", "20.50"), (second["id"], "Live", "10.25")]
    workbook = load_workbook(io.BytesIO(client.get(f"/api/reports/{report['id']}/export/xlsx").content), data_only=True)
    source_labels = [cell.value for sheet in workbook for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("Source:")]
    assert any(" / History" in label for label in source_labels)
    assert any(" / Live" in label for label in source_labels)


def test_ungranted_extra_source_stops_report_instead_of_omitting_requested_analysis(client, admin):
    subscribe(client)
    first, _ = fixture_source(client)
    result = client.post("/api/reports", headers=HEADERS, json={"title": "Unavailable extra source", "source_id": first["id"], "query": query_plan(), "analyses": [{"source_id": "not-granted", "query": query_plan()}]}).json()
    assert wait_report(client, result["id"])["status"] == "failed"
    assert client.get(f"/api/reports/{result['id']}/export/xlsx").status_code != 200


def test_simultaneous_queries_keep_their_selected_databases_separate(client, admin):
    source, _ = fixture_source(client)
    state = client.app.state
    original = MultipleDatabaseFixture(replace(state.settings, sqlserver_database="Live"), state.connectors.catalog)
    state.connectors.register(source["id"], original)

    async def query_both():
        return await asyncio.gather(*(state.connectors.database_call(admin["id"], source["id"], "report_query", database=database, query=query_plan()) for database in ("Live", "History")))

    results = client.portal.call(query_both)
    assert [(r["database"], r["rows"][0]["revenue"]) for r in results] == [("Live", "10.25"), ("History", "20.50")]
    assert original.settings.sqlserver_database == "Live"


def test_all_database_tools_expose_target_and_guidance_requires_full_inventory():
    tools = {tool["name"]: tool for tool in REPORT_TOOLS[0]["tools"]}
    for name in ("discover_sales", "list_tables", "describe_table", "query_database", "create_database_report"):
        assert "database" in tools[name]["inputSchema"]["properties"]
    assert "discover_databases" in tools and "list_databases" in tools
    assert "not just the configured/default database" in BASE_INSTRUCTIONS
