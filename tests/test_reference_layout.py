import io
import json

import pytest
import xlsxwriter
from openpyxl import load_workbook

from app.attachment_reader import extract
from app.config import Settings
from app.provider import BASE_INSTRUCTIONS
from app.reports import WORKBOOK_SCRIPT_API
from app.analysis_explanation import ANALYSIS_EXPLANATION
from app.workbook_runner import WorkbookRunner
from app.workbook_script import run_script
from app.workbooks import ReportDataError
from test_excel_reports import workbook_fixture
from conftest import HEADERS


def reference_bytes():
    data=io.BytesIO()
    book=xlsxwriter.Workbook(data)
    s=book.add_worksheet('Performance')
    header=book.add_format({'bold':True,'bg_color':'#17365D','font_color':'#FFFFFF','align':'center'})
    body=book.add_format({'num_format':'#,##0.00'})
    percent=book.add_format({'num_format':'0.00%'})
    total=book.add_format({'bold':True,'bg_color':'#DCE6F1','num_format':'#,##0.00'})
    s.merge_range('B1:F1','Historical report 18 October',header)
    s.merge_range('B4:B6','Branch',header)
    s.merge_range('C4:D4','DAY 18.10',header)
    s.merge_range('E4:F4','MONTH 1–18',header)
    for c,label in enumerate(['Sales','Margin','Sales','Margin'],2):
        s.write(4,c,label,header);s.write(5,c,'2026',header)
    for r,label in [(6,'Branch A'),(7,'TOTAL')]:
        s.write(r,1,label,total if r==7 else body)
        for c in range(2,6):s.write_number(r,c,999999,total if r==7 else percent if c in [3,5] else body)
    s.set_column('A:A',3);s.set_column('B:B',28);s.set_column('C:F',14)
    s.freeze_panes('C7');s.set_landscape()
    book.close();return data.getvalue()


def prepared(tmp_path,script=None):
    path=tmp_path/'reference.xlsx';path.write_bytes(reference_bytes())
    reference=extract(path,'.xlsx')['format']
    rows=[{'name':'Branch A','day_sales':'100','day_margin':'0.25','month_sales':'1000','month_margin':'0.30'},
          {'name':'TOTAL','day_sales':'100','day_margin':'0.25','month_sales':'1000','month_margin':None}]
    report,snapshot=workbook_fixture(rows)
    snapshot['provenance_json']=json.dumps({'columns':[{'name':k,'type':'text' if k=='name' else 'decimal'} for k in rows[0]],'format_reference':reference})
    document=json.loads(report['content_json'])
    headers=[['Pika','DITA 17.10',None,'MUAJI 1–17',None],[None,'Shitjet','Marzhi','Shitjet','Marzhi'],[None,'2026','2026','2026','2026']]
    script=script or "def build(report):\n    s=report.sheet('Performanca')\n    s.title('Raporti i 17 tetorit')\n    s.reference_table(0, 'Performance', "+repr(headers)+", formats={'day_margin':'percent','month_margin':'percent'})\n"
    document['workbook']={'language':'sq','labels':[],'analyses':[{'index':0,'name':'Performanca','scope':'Të gjithë rreshtat','query_explanation':'Të dhëna prove','columns':[{'column':k,'label':k} for k in rows[0]]}],'python_script':script}
    report['content_json']=json.dumps(document)
    return report,snapshot,document


@pytest.mark.asyncio
async def test_reference_preserves_grouped_headers_percentages_subtotals_and_all_values(tmp_path):
    report,snapshot,document=prepared(tmp_path)
    content,checks=await WorkbookRunner(Settings(data_dir=tmp_path,env_file=None)).build(report,snapshot)
    assert checks['source_rows']==[2]
    book=load_workbook(io.BytesIO(content),data_only=True)
    s=book['Performanca']
    assert s['B4'].value=='Pika' and s['C4'].value=='DITA 17.10'
    assert 'B4:B6' in s.merged_cells and 'C4:D4' in s.merged_cells
    assert s['C7'].value==100 and s['D7'].value==.25 and s['D7'].number_format.startswith('0.00%')
    assert s['F8'].value is None
    assert s['B8'].font.bold and s['B8'].fill.fgColor.rgb.endswith('DCE6F1')
    assert s.freeze_panes=='C7' and s.print_area and s.page_setup.orientation=='landscape'
    assert not any(c.value==999999 for row in s for c in row)
    assert not any('18.10' in str(c.value) for row in s for c in row)
    assert s.column_dimensions['B'].width>=28


def test_flattening_a_compatible_grouped_reference_is_rejected(tmp_path):
    report,snapshot,doc=prepared(tmp_path,"def build(report):\n    s=report.sheet('Output')\n    s.table(0)\n")
    with pytest.raises(ReportDataError,match='reference_table'):
        run_script(doc['workbook']['python_script'],doc,snapshot)


@pytest.mark.asyncio
async def test_reference_header_merges_cannot_hide_labels(tmp_path):
    report,snapshot,doc=prepared(tmp_path)
    script=doc['workbook']['python_script'].replace("'DITA 17.10', None","'DITA 17.10', 'hidden label'")
    doc['workbook']['python_script']=script;report['content_json']=json.dumps(doc)
    with pytest.raises(ReportDataError,match='hide another label'):
        await WorkbookRunner(Settings(data_dir=tmp_path,env_file=None)).build(report,snapshot)


def test_existing_upload_metadata_upgrades_without_touching_data_or_grants(client,admin):
    cid=client.post('/api/conversations',headers=HEADERS,json={}).json()['id']
    item=client.post(f'/api/conversations/{cid}/attachments',headers=HEADERS,files={'file':('reference.xlsx',reference_bytes())}).json()
    store=client.app.state.attachments
    row=store.owned(admin['id'],item['id'],cid)
    before=store.parsed(row)
    legacy={**before,'format':{k:v for k,v in before['format'].items() if k not in ['layout_version','sheet_layouts','styles','layout_policy']}}
    (store.directory/(item['id']+'.json')).write_text(json.dumps(legacy))
    client.portal.call(store.ensure_format_reference,admin['id'],item['id'],cid)
    after=store.parsed(row)
    assert after['format']['layout_version']==2 and after['tables']==before['tables'] and after['text']==before['text']
    assert store.format_reference(admin['id'],item['id'],cid)['sheet_layouts'][0]['name']=='Performance'
    from fastapi import HTTPException
    with pytest.raises(HTTPException):
        client.portal.call(store.ensure_format_reference,'another-user',item['id'],cid)


def test_shared_analysis_policy_explains_scope_method_checks_and_limits():
    assert ANALYSIS_EXPLANATION in BASE_INSTRUCTIONS
    assert 's.reference_table' in WORKBOOK_SCRIPT_API and 'formats maps actual fields' in WORKBOOK_SCRIPT_API
    for phrase in ['Scope and inputs','Inclusion and exclusions','Method:','Findings and checks','Limits and next steps','private model reasoning']:
        assert phrase in ANALYSIS_EXPLANATION


@pytest.mark.asyncio
async def test_plain_output_table_formats_ratios_without_changing_the_numbers(tmp_path):
    report,snapshot,doc=prepared(tmp_path)
    provenance=json.loads(snapshot['provenance_json']);provenance.pop('format_reference')
    snapshot['provenance_json']=json.dumps(provenance)
    doc['workbook']['python_script']="def build(report):\n    s=report.sheet('Prova')\n    s.table(0,formats={'day_margin':'percent','month_margin':'percent'})\n"
    report['content_json']=json.dumps(doc)
    content,_=await WorkbookRunner(Settings(data_dir=tmp_path,env_file=None)).build(report,snapshot)
    sheet=load_workbook(io.BytesIO(content),data_only=True)['Prova']
    assert sheet['C3'].value==.25 and sheet['C3'].number_format.startswith('0.00%')


@pytest.fixture
def claude(client,monkeypatch):
    from test_claude_subscriptions import claude as fixture
    return fixture.__wrapped__(client,monkeypatch)


@pytest.mark.parametrize('provider',['chatgpt','claude'])
def test_background_analysis_explanation_reaches_both_providers(client,admin,claude,provider):
    from test_chat_controls import prepare
    from test_excel_reports import wait_report
    cid=prepare(client,provider)
    created=client.post('/api/reports',headers=HEADERS,json={'title':'Analizë prove','conversation_id':cid,'rows':[{'value':10}]}).json()
    assert wait_report(client,created['id'])['status']=='completed'
    if provider=='claude':
        prompt=json.loads(next(c.content for c in reversed(claude) if c.options.output_format))
        assert ANALYSIS_EXPLANATION in claude[-1].options.system_prompt
    else:
        path=client.app.state.settings.data_dir/'codex'/admin['id']/'last-turn-input.json'
        prompt=json.loads(json.loads(path.read_text())[0]['text'])
    assert prompt['analysis_explanation_requirements']==ANALYSIS_EXPLANATION
    assert 'reference_table' in prompt['workbook_script_api']


def test_default_layout_cannot_flatten_a_selected_grouped_reference(tmp_path):
    report,snapshot,doc=prepared(tmp_path,"def build(report):\n    report.default_layout()\n")
    with pytest.raises(ReportDataError,match='reference_table'):
        run_script(doc['workbook']['python_script'],doc,snapshot)
