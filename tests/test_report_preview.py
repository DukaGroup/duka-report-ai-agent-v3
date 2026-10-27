from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import pytest
import xlsxwriter

from app.connectors import SourceNotConfigured
from app.workbook_preview import PREVIEW_COLUMNS, PREVIEW_ROWS, PreviewSheetNotFound, workbook_preview
from conftest import HEADERS
from test_application import create_member, sign_in_member
from test_excel_reports import subscribe, wait_report


def test_workbook_preview_pages_through_every_sheet_row_and_column(tmp_path):
    path = tmp_path / "report.xlsx"
    book = Workbook()
    first = book.active
    first.title = "Përmbledhje"
    first["A1"] = "Titull " + "i plotë " * 80
    first["A1"].font = Font(bold=True, color="FFFFFFFF")
    first["A1"].fill = PatternFill("solid", fgColor="FF205C49")
    first.merge_cells("A1:C1")
    first["B2"] = 123
    first["B2"].number_format = "#,##0.00"
    first.cell(row=PREVIEW_ROWS + 5, column=PREVIEW_COLUMNS + 4, value="Last cell")
    second = book.create_sheet("Detaje")
    second["A1"] = "Second sheet"
    book.save(path)

    preview = workbook_preview(path)
    assert [sheet["name"] for sheet in preview["sheets"]] == ["Përmbledhje", "Detaje"]
    assert preview["visible_rows"] == PREVIEW_ROWS
    assert preview["visible_columns"] == PREVIEW_COLUMNS
    assert preview["has_more_rows"] and preview["has_more_columns"]
    assert preview["rows"][0][0]["text"] == first["A1"].value
    assert preview["rows"][0][0]["bold"]
    assert preview["rows"][0][0]["color"] == "#FFFFFF"
    assert preview["rows"][0][0]["background"] == "#205C49"
    assert preview["rows"][1][1]["text"] == "123.00"
    assert preview["rows"][1][1]["raw"] == "123"
    assert preview["merged"] == [{"start_row": 1, "start_column": 1, "end_row": 1, "end_column": 3}]
    assert len(preview["rows"]) == PREVIEW_ROWS
    assert all(len(row) == PREVIEW_COLUMNS for row in preview["rows"])
    last = workbook_preview(path, 0, PREVIEW_ROWS + 1, PREVIEW_COLUMNS + 1)
    assert last["row_start"] == PREVIEW_ROWS + 1
    assert last["column_start"] == PREVIEW_COLUMNS + 1
    assert last["rows"][4][3]["text"] == "Last cell"
    assert not last["has_more_rows"] and not last["has_more_columns"]
    assert workbook_preview(path, 1)["rows"][0][0]["text"] == "Second sheet"
    with pytest.raises(PreviewSheetNotFound):
        workbook_preview(path, 2)


def test_workbook_preview_includes_saved_chart_values(tmp_path):
    path = tmp_path / "chart.xlsx"
    book = xlsxwriter.Workbook(str(path))
    sheet = book.add_worksheet("Grafiku")
    sheet.write_row("A1", ["Muaji", "Shitjet"])
    sheet.write_row("A2", ["Janar", 12])
    sheet.write_row("A3", ["Shkurt", 18])
    chart = book.add_chart({"type": "column"})
    chart.add_series({
        "name": ["Grafiku", 0, 1],
        "categories": ["Grafiku", 1, 0, 2, 0],
        "values": ["Grafiku", 1, 1, 2, 1],
    })
    chart.set_title({"name": "Shitjet sipas muajit"})
    sheet.insert_chart("D2", chart)
    book.close()

    preview = workbook_preview(path)
    assert preview["charts"] == [{
        "title": "Shitjet sipas muajit",
        "type": "column",
        "categories": ["Janar", "Shkurt"],
        "series": [{"name": "Shitjet", "values": [12.0, 18.0]}],
        "truncated": False,
    }]


def test_report_preview_requires_current_access_and_preserves_saved_excel(client, admin, monkeypatch):
    subscribe(client)
    created = client.post(
        "/api/reports",
        headers=HEADERS,
        json={"title": "Preview before download", "rows": [{"revenue": 120}]},
    ).json()
    report = wait_report(client, created["id"])
    assert report["status"] == "completed"
    path = client.app.state.settings.data_dir / "reports" / f"{report['id']}.xlsx"
    before = path.read_bytes()

    response = client.get(f"/api/reports/{report['id']}/preview")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    preview = response.json()
    assert preview["sheets"] and preview["rows"]
    assert preview["sheet_index"] == 0
    assert preview["row_start"] == 1 and preview["column_start"] == 1
    assert path.read_bytes() == before
    last_cell = client.get(f"/api/reports/{report['id']}/preview?row=2&column=2")
    assert last_cell.status_code == 200
    assert last_cell.json()["row_start"] == 2
    assert last_cell.json()["column_start"] == 2
    assert client.get(f"/api/reports/{report['id']}/preview?sheet=999").status_code == 404
    assert client.get(f"/api/reports/{report['id']}/preview?sheet=-1").status_code == 422
    client.app.state.db.execute("UPDATE reports SET status='running' WHERE id=?", (report["id"],))
    assert client.get(f"/api/reports/{report['id']}/preview").status_code == 409
    client.app.state.db.execute("UPDATE reports SET status='completed' WHERE id=?", (report["id"],))

    def revoke_access(*_args, **_kwargs):
        raise SourceNotConfigured("revoked")

    with monkeypatch.context() as patch:
        patch.setattr("app.main.authorize_saved_report", revoke_access)
        assert client.get(f"/api/reports/{report['id']}/preview").status_code == 403

    create_member(client)
    sign_in_member(client)
    assert client.get(f"/api/reports/{report['id']}/preview").status_code == 404
