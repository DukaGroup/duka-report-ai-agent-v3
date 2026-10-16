import asyncio
import io
import json
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.config import Settings
from app.provider import CodexRuntime
from app.workbooks import ReportDataError, excel_report
from app.workbook_runner import WorkbookRunner
from app.workbook_script import validate_script
from app.workbook_validation import validate_workbook
from conftest import HEADERS
from test_excel_reports import workbook_fixture, subscribe, wait_report


def scripted_fixture(script, rows=None):
    report, snapshot = workbook_fixture(rows)
    document = json.loads(report["content_json"])
    document["workbook"] = {"language": "en", "labels": [], "python_script": script,
                            "analyses": [{"index": 0, "name": "Monthly sales", "scope": "All rows", "query_explanation": "Fixture",
                                          "columns": [{"column": "month", "label": "Month"}, {"column": "revenue", "label": "Revenue"}, {"column": "cost", "label": "Cost"}]}]}
    report["content_json"] = json.dumps(document)
    return report, snapshot, document


@pytest.mark.parametrize("count", [1, 21, 100])
async def test_model_script_can_size_every_supported_column_without_dropping_data(tmp_path, count):
    rows = [{f"field{i}": f"record{r}-column{i}" for i in range(count)} for r in range(3)]
    report, snapshot, document = scripted_fixture(
        "def build(report):\n    s=report.sheet('Wide detail')\n    s.widths(" + str([18] * count) + ")\n    s.table(0)\n", rows)
    columns = [{"name": field, "type": "text"} for field in rows[0]]
    snapshot["provenance_json"] = json.dumps({"source": "Fixture", "columns": columns, "analysis": {"kind": "data"}})
    document["workbook"]["analyses"][0]["columns"] = [{"column": field, "label": field} for field in rows[0]]
    report["content_json"] = json.dumps(document)
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    content, verified = await WorkbookRunner(settings).build(report, snapshot)
    assert verified["source_rows"] == [3]
    book = load_workbook(io.BytesIO(content), data_only=True)
    values = {cell.value for row in book["Wide detail"] for cell in row}
    assert all(value in values for row in rows for value in row.values())


@pytest.mark.asyncio
async def test_generated_python_builds_custom_workbook_and_every_source_row_is_verified(tmp_path):
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    rows = [{"month": f"2026-01-{(i % 28)+1:02d}", "revenue": i, "cost": 1} for i in range(1200)]
    report, snapshot, _ = scripted_fixture("""def build(report):
    s = report.sheet('Sales review')
    s.title()
    s.metric(0, 'revenue', 'Sales', currency='ALL')
    s.chart(0, 'month', ['revenue'], 'Sales trend', chart_type='line', currency='ALL')
    s.table(0)
    s.narrative()
""", rows)
    content, verified = await WorkbookRunner(settings).build(report, snapshot)
    assert verified == {"source_rows": [1200], "verified_calculations": 1}
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    assert workbook.sheetnames[0] == "Sales review"
    assert workbook.worksheets[0]["C4"].value == sum(range(1200))
    assert len(workbook.worksheets[0]._charts) == 1
    source = workbook["1 Monthly sales"]
    assert len(list(source.iter_rows(min_row=6, max_row=1205, max_col=3))) == 1200
    # Date sorting still keeps every row's matching amount/cost together.
    assert sorted(source.cell(row, 2).value for row in range(6, 1206)) == list(range(1200))
    assert list((tmp_path / "excel-jobs").iterdir()) == []


@pytest.mark.parametrize("script", [
    "import os\ndef build(report): pass",
    "def build(report):\n    open('/run/duka.env')",
    "def build(report):\n    x = report.__class__",
    "def build(report):\n    x = getattr(report, 'sheet')",
    "def build(report):\n    report.sheet = str",
    "def build(report):\n    x = '{0.__class__}'.format(report)",
    "def build(report):\n    exec('print(1)')",
])
def test_unapproved_generated_python_is_rejected_before_execution(script):
    with pytest.raises(ReportDataError):
        validate_script(script)


def test_kernel_sandbox_blocks_file_opens_network_and_child_processes(tmp_path):
    job = tmp_path / "job"
    job.mkdir()
    outside = tmp_path / "private-secret"
    outside.write_text("fixture credential")
    code = """import os,socket,sys
from app.script_sandbox import isolate
isolate(sys.argv[1])
for operation in [lambda: os.open(sys.argv[2],os.O_RDONLY),lambda: socket.socket(),lambda: os.fork()]:
    try:
        operation()
    except PermissionError:
        continue
    raise AssertionError('restricted operation unexpectedly succeeded')
print('all denied')
"""
    result = subprocess.run([sys.executable, "-c", code, str(job), str(outside)],
                            env={"PYTHONPATH": str(Path("backend").resolve())}, capture_output=True, timeout=10)
    assert result.returncode == 0, result.stderr.decode()
    assert result.stdout.strip() == b"all denied"


@pytest.mark.asyncio
async def test_time_limit_kills_worker_and_cleans_private_job(tmp_path):
    settings = Settings(data_dir=tmp_path, env_file=None, workbook_timeout=1)
    settings.prepare()
    report, snapshot, _ = scripted_fixture("def build(report):\n    for i in range(1000000000):\n        x = i + 1\n    report.default_layout()\n")
    with pytest.raises(ReportDataError, match="limit|stopped"):
        await WorkbookRunner(settings).build(report, snapshot)
    assert list((tmp_path / "excel-jobs").iterdir()) == []


@pytest.mark.asyncio
async def test_cancellation_terminates_worker_and_leaves_no_artifact(tmp_path):
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    report, snapshot, _ = scripted_fixture("def build(report):\n    for i in range(1000000000):\n        x = i + 1\n    report.default_layout()\n")
    task = asyncio.create_task(WorkbookRunner(settings).build(report, snapshot))
    await asyncio.sleep(0.3)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert list((tmp_path / "excel-jobs").iterdir()) == []


@pytest.mark.asyncio
async def test_another_workbook_finishes_while_one_script_is_still_running(tmp_path):
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    slow, snapshot, _ = scripted_fixture("def build(report):\n    for i in range(1000000000):\n        x = i + 1\n    report.default_layout()\n")
    fast, _, _ = scripted_fixture("def build(report):\n    report.default_layout()\n")
    runner = WorkbookRunner(settings)
    first = asyncio.create_task(runner.build(slow, snapshot))
    content, _ = await asyncio.wait_for(runner.build(fast, snapshot), 5)
    assert content.startswith(b"PK") and not first.done()
    first.cancel()
    with pytest.raises(asyncio.CancelledError):
        await first


@pytest.mark.asyncio
async def test_generated_calculations_keep_all_and_foreign_currencies_separate(tmp_path):
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    rows = [{"month": "2026-01-01", "revenue": 10, "cost": 2, "Currency Code": ""},
            {"month": "2026-01-02", "revenue": 20, "cost": 3, "Currency Code": "EUR"}]
    source = "def build(report):\n    s = report.sheet('Currency totals')\n    s.title()\n    for currency in report.currencies(0):\n        s.metric(0, 'revenue', 'Sales', currency=currency)\n"
    report, snapshot, document = scripted_fixture(source, rows)
    document["workbook"]["analyses"][0]["columns"].append({"column": "Currency Code", "label": "Currency"})
    report["content_json"] = json.dumps(document)
    provenance = json.loads(snapshot["provenance_json"])
    provenance["columns"].append({"name": "Currency Code", "type": "nvarchar"})
    snapshot["provenance_json"] = json.dumps(provenance)
    content, verified = await WorkbookRunner(settings).build(report, snapshot)
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    assert workbook.worksheets[0]["C4"].value == 10
    assert workbook.worksheets[0]["C6"].value == 20
    assert verified["verified_calculations"] == 2
    document["workbook"]["python_script"] = "def build(report):\n    s = report.sheet('Bad total')\n    s.metric(0, 'revenue', 'Sales')\n"
    report["content_json"] = json.dumps(document)
    with pytest.raises(ReportDataError, match="incompatible currencies"):
        await WorkbookRunner(settings).build(report, snapshot)


@pytest.mark.asyncio
async def test_custom_tables_preserve_numbers_beyond_excel_precision_as_exact_text(tmp_path):
    settings = Settings(data_dir=tmp_path, env_file=None)
    settings.prepare()
    exact = 100000000000000000001
    report, snapshot, _ = scripted_fixture("def build(report):\n    s = report.sheet('Exact values')\n    s.table(0)\n",
                                          [{"month": "2026-01-01", "revenue": exact, "cost": 1}])
    content, verified = await WorkbookRunner(settings).build(report, snapshot)
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    assert workbook.worksheets[0]["B3"].value == str(exact)
    assert workbook["1 Monthly sales"]["B6"].value == str(exact)
    assert verified["source_rows"] == [1]


def test_source_cell_tampering_and_shifted_rows_cannot_pass_validation():
    report, snapshot, document = scripted_fixture("def build(report):\n    report.default_layout()\n")
    content = excel_report(report, snapshot)
    validate_workbook(content, document, snapshot, [])
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(buffer, "w") as modified:
        for name in original.namelist():
            value = original.read(name)
            if name == "xl/worksheets/sheet2.xml":
                value = value.replace(b"<v>100</v>", b"<v>999</v>")
            modified.writestr(name, value)
    with pytest.raises(ReportDataError, match="source values"):
        validate_workbook(buffer.getvalue(), document, snapshot, [])


@pytest.mark.asyncio
async def test_ooxml_control_characters_and_literal_escape_strings_validate_without_hiding_changes(tmp_path):
    settings = Settings(data_dir=tmp_path,env_file=None)
    settings.prepare()
    strings = ['Line one\r\nLine two','Control\x01text','Literal _x000D_ and _x005F_','Emoji 😊']
    rows = [{'month':text,'revenue':i+1,'cost':1} for i,text in enumerate(strings)]
    report,snapshot,document = scripted_fixture("def build(report):\n    s = report.sheet('Source text')\n    s.table(0)\n",rows)
    provenance = json.loads(snapshot['provenance_json'])
    provenance['options'] = {'kind':'data'}
    provenance['columns'][0]['type'] = 'text'
    snapshot['provenance_json'] = json.dumps(provenance)
    content,verified = await WorkbookRunner(settings).build(report,snapshot)
    assert verified['source_rows'] == [4]
    changed = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(changed,'w') as modified:
        for name in original.namelist():
            value = original.read(name)
            if name == 'xl/sharedStrings.xml':
                value = value.replace(b'Line one_x000D_',b'Line one_x005F_x000D_')
            modified.writestr(name,value)
    with pytest.raises(ReportDataError,match='source values|misaligned'):
        validate_workbook(changed.getvalue(),document,snapshot,[],True)


def test_bad_script_is_repaired_once_and_only_verified_workbook_is_published(client, admin, monkeypatch):
    subscribe(client)
    original = CodexRuntime.stream
    attempts = []

    async def scripted(self, thread_id, content, model=None, output_schema=None):
        if output_schema:
            attempts.append(json.loads(content))
        async for event in original(self, thread_id, content, model, output_schema):
            if output_schema and len(attempts) == 1 and event["type"] == "done":
                doc = json.loads(event["text"])
                doc["workbook"]["python_script"] = "def build(report):\n    open('/run/duka.env')\n"
                event = {**event, "text": json.dumps(doc)}
            yield event

    monkeypatch.setattr(CodexRuntime, "stream", scripted)
    created = client.post("/api/reports", headers=HEADERS, json={"title": "Repair", "rows": [{"revenue": 42}]}).json()
    report = wait_report(client, created["id"])
    assert report["status"] == "completed", report
    assert len(attempts) == 2 and "validation_error" in attempts[1]
    assert client.get(f"/api/reports/{created['id']}/export/xlsx").status_code == 200
    assert "default_layout" in report["document"]["workbook"]["python_script"]


def test_failed_repairs_never_publish_a_partial_file(client, admin, monkeypatch):
    subscribe(client)
    original = CodexRuntime.stream
    attempts = []

    async def invalid(self, thread_id, content, model=None, output_schema=None):
        if output_schema:
            attempts.append(content)
        async for event in original(self, thread_id, content, model, output_schema):
            if output_schema and event["type"] == "done":
                doc = json.loads(event["text"])
                target = doc["workbook"] if "workbook" in doc else doc
                target["python_script"] = "def build(report):\n    open('/run/duka.env')\n"
                event = {**event, "text": json.dumps(doc)}
            yield event

    monkeypatch.setattr(CodexRuntime, "stream", invalid)
    created = client.post("/api/reports", headers=HEADERS, json={"title": "Reject", "rows": [{"revenue": 42}]}).json()
    report = wait_report(client, created["id"])
    assert report["status"] == "failed" and len(attempts) == 3
    assert client.get(f"/api/reports/{created['id']}/export/xlsx").status_code == 409
    assert not (client.app.state.settings.data_dir / "reports" / (created["id"]+".xlsx")).exists()
